"""Regression: runs the WIRED production V3 pipeline (not the standalone
step1/step2 in-memory harness) against the regression contract end to end —
real Document/DocumentPage rows, real DB persistence via
`run_and_persist_v3_extraction` — then reports dataset row counts and checks
the confirmed-bad screenshot values never reappear as business fields.

Usage: python backend/scripts/generate_v3_workbook.py
"""

from __future__ import annotations

import hashlib
import shutil
import sys
import uuid
from collections import Counter
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from sqlalchemy import delete, select  # noqa: E402

from app.core.config import get_settings  # noqa: E402
from app.database.session import SessionLocal  # noqa: E402
from app.models.document import Document  # noqa: E402
from app.models.document_attachment import DocumentAttachment  # noqa: E402
from app.models.document_clause_reference import DocumentClauseReference  # noqa: E402
from app.models.document_contract_summary import DocumentContractSummary  # noqa: E402
from app.models.document_funding_line import DocumentFundingLine  # noqa: E402
from app.models.document_line_item import DocumentLineItem  # noqa: E402
from app.models.document_metadata_field import DocumentMetadataField  # noqa: E402
from app.models.document_page import DocumentPage  # noqa: E402
from app.models.document_performance_period import DocumentPerformancePeriod  # noqa: E402
from app.models.document_qa_review import DocumentQaReview  # noqa: E402
from app.services.document_extraction import process_document_pages  # noqa: E402
from app.services.v3_export_builder import build_v3_xlsx  # noqa: E402
from app.services.v3_orchestrator import run_and_persist_v3_extraction  # noqa: E402
from app.services.v3_reader import get_normalized_v3_document  # noqa: E402

PDF_PATH = BACKEND_ROOT.parent / "reference" / "regression" / "Contract_47QRCA25DSF07 (2).pdf"
GROUND_TRUTH_XLSX = BACKEND_ROOT.parent / "reference" / "schemas" / "V3_Extraction.xlsx"
GENERATED_XLSX_PATH = BACKEND_ROOT.parent / "reference" / "regression" / "generated_v3.xlsx"

KNOWN_BAD_EXAMPLES = [
    ("Amount", "$0.00"),
    ("NAICS", "Codes"),
    ("Solicitation Number", "whether or not Contractors"),
    ("Adm 4800", None),
    ("B.8", "1 CONUS Standardized Labor Categories"),
    ("C.1", "1 North American Industry Classification System"),
    ("F.4", "1 Deliverable and Reporting Requirements"),
    ("F.5", None),
    ("FAR", "TITLE and DATE"),
]

MARKER_FILENAME = "__v3_regression_marker__.pdf"


def _get_or_create_document(database) -> Document:
    settings = get_settings()
    checksum = hashlib.sha256(PDF_PATH.read_bytes()).hexdigest()

    existing = database.scalars(
        select(Document).where(Document.checksum_sha256 == checksum)
    ).first()
    if existing is not None:
        return existing

    document_id = str(uuid.uuid4())
    stored_filename = f"{document_id}.pdf"
    destination = settings.upload_path / stored_filename
    settings.upload_path.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(PDF_PATH, destination)

    import fitz

    with fitz.open(str(destination)) as doc:
        page_count = doc.page_count

    document = Document(
        id=document_id,
        original_filename=PDF_PATH.name,
        stored_filename=stored_filename,
        content_type="application/pdf",
        size_bytes=destination.stat().st_size,
        checksum_sha256=checksum,
        page_count=page_count,
        encrypted=False,
        status="ready",
    )
    database.add(document)
    database.commit()
    return document


def main() -> None:
    database = SessionLocal()
    settings = get_settings()

    document = _get_or_create_document(database)
    print(f"Document: {document.id} ({document.original_filename}, {document.page_count} pages)")

    pages = list(
        database.scalars(
            select(DocumentPage)
            .where(DocumentPage.document_id == document.id)
            .order_by(DocumentPage.page_number)
        )
    )
    if not pages:
        print("No DocumentPage rows yet — running process_document_pages (native text only)...")
        file_path = settings.upload_path / document.stored_filename
        process_document_pages(
            database=database,
            document_record=document,
            file_path=file_path,
            settings=settings,
            run_ocr=False,
            page_start=None,
            page_end=None,
            force_reprocess=False,
        )
        database.commit()
        pages = list(
            database.scalars(
                select(DocumentPage)
                .where(DocumentPage.document_id == document.id)
                .order_by(DocumentPage.page_number)
            )
        )
    print(f"DocumentPage rows: {len(pages)}")

    print("\nRunning run_and_persist_v3_extraction (the real production V3 stage)...")
    summary = run_and_persist_v3_extraction(database=database, document=document, pages=pages)

    print("\n" + "=" * 100)
    print("V3 EXTRACTION SUMMARY (wired production pipeline)")
    print("=" * 100)
    print(f"Candidates classified: {summary.candidate_count}")
    print(f"Pages classified: {summary.pages_classified} | with geometry: {summary.pages_with_geometry}")
    print(f"All Fields rows: {summary.all_fields_count}")
    print(f"CLIN rows: {summary.clin_count}")
    print(f"Funding rows: {summary.funding_count}")
    print(f"Performance/Delivery rows: {summary.performance_delivery_count}")
    print(f"Attachment rows: {summary.attachment_count}")
    print(f"Clause reference rows (Clauses+DFARS+FAR References combined): {summary.clause_reference_count}")
    if summary.warnings:
        print("Warnings:")
        for warning in summary.warnings:
            print(f"  - {warning}")

    all_fields = list(
        database.scalars(
            select(DocumentMetadataField).where(
                DocumentMetadataField.document_id == document.id,
                DocumentMetadataField.extraction_source == "v3",
            )
        )
    )
    clins = list(
        database.scalars(
            select(DocumentLineItem).where(DocumentLineItem.document_id == document.id)
        )
    )
    clauses = list(
        database.scalars(
            select(DocumentClauseReference).where(
                DocumentClauseReference.document_id == document.id
            )
        )
    )
    contract_summary = database.scalars(
        select(DocumentContractSummary).where(
            DocumentContractSummary.document_id == document.id
        )
    ).first()
    qa_review = list(
        database.scalars(
            select(DocumentQaReview).where(DocumentQaReview.document_id == document.id)
        )
    )

    print("\n" + "=" * 100)
    print("KNOWN-BAD-RECORD REGRESSION (must all be 0 as GENERAL_ACCEPTED_FIELD/All Fields)")
    print("=" * 100)
    bad_count = 0
    for old_field, old_value in KNOWN_BAD_EXAMPLES:
        survivors = [
            f
            for f in all_fields
            if old_field.lower() in (f.label or "").lower()
            and (old_value is None or old_value.lower() in (f.value or "").lower())
        ]
        status = "PASS (excluded)" if not survivors else "FAIL (reappeared!)"
        if survivors:
            bad_count += len(survivors)
        print(f"  {old_field!r} -> {old_value!r}: {status}")
    print(f"\nTotal known-bad-record survivors in All Fields: {bad_count} (must be 0)")

    kv_visible = [f for f in all_fields if (f.field_key or "").startswith("kv_")]
    print(f"Raw kv_* visible in V3 All Fields: {len(kv_visible)} (must be 0)")

    print("\n" + "=" * 100)
    print("CONTRACT SUMMARY")
    print("=" * 100)
    if contract_summary is None:
        print("  No Contract Summary row produced.")
    else:
        for column in (
            "contract_number", "solicitation_rfp", "contract_vehicle", "agency_office",
            "contractor", "award_date", "ceiling_max_aggregate", "minimum_guarantee",
            "base_period", "options", "max_duration", "task_order_range", "naics",
            "size_standard", "qa_status",
        ):
            print(f"  {column}: {getattr(contract_summary, column)!r}")

    print("\n" + "=" * 100)
    print("CLAUSE FAMILY / CONTEXT BREAKDOWN")
    print("=" * 100)
    breakdown = Counter((c.clause_family, c.citation_context) for c in clauses)
    for (family, context), count in sorted(breakdown.items()):
        print(f"  {family} / {context}: {count}")

    print("\n" + "=" * 100)
    print("QA REVIEW ROLLUP")
    print("=" * 100)
    for row in qa_review:
        print(f"  {row.qa_check}: {row.result} — {row.details} (action: {row.action})")

    print("\n" + "=" * 100)
    print("STRUCTURAL COMPARISON vs reference/schemas/V3_Extraction.xlsx")
    print("=" * 100)
    import openpyxl

    normalized = get_normalized_v3_document(database, document.id)
    generated_bytes = build_v3_xlsx(normalized)
    GENERATED_XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    GENERATED_XLSX_PATH.write_bytes(generated_bytes)
    print(f"Wrote {GENERATED_XLSX_PATH}")

    import io

    generated_wb = openpyxl.load_workbook(io.BytesIO(generated_bytes))
    ground_truth_wb = openpyxl.load_workbook(str(GROUND_TRUTH_XLSX), read_only=True)

    sheets_match = generated_wb.sheetnames == ground_truth_wb.sheetnames
    print(f"Sheet names + order match: {sheets_match}")
    if not sheets_match:
        print(f"  generated: {generated_wb.sheetnames}")
        print(f"  ground truth: {ground_truth_wb.sheetnames}")

    all_headers_match = True
    for name in ground_truth_wb.sheetnames:
        if name == "README":
            continue
        gt_headers = [c.value for c in next(ground_truth_wb[name].iter_rows(min_row=3, max_row=3))]
        gen_headers = [c.value for c in next(generated_wb[name].iter_rows(min_row=1, max_row=1))]
        match = gt_headers == gen_headers
        all_headers_match = all_headers_match and match
        print(f"  {name}: header match={match}")
        if not match:
            print(f"    GT : {gt_headers}")
            print(f"    GEN: {gen_headers}")

    print(f"\nOVERALL STRUCTURAL MATCH: {sheets_match and all_headers_match}")

    print("\n" + "=" * 100)
    print("DATASET CONSISTENCY — normalized reader vs generated XLSX (single canonical reader)")
    print("=" * 100)
    xlsx_sheet_by_dataset = {
        "clins": "CLINs",
        "funding": "Funding",
        "performance_delivery": "Performance Delivery",
        "attachments": "Attachments",
        "clauses": "Clauses",
        "far_references": "FAR References",
        "dfars": "DFARS",
        "all_fields": "All Fields",
        "qa_review": "QA Review",
        "source_documents": "Source Documents",
    }
    print(f"{'DATASET':<22} {'NORMALIZED':>10} {'XLSX':>8}  MATCH")
    all_match = True
    for attr, sheet_name in xlsx_sheet_by_dataset.items():
        normalized_count = len(getattr(normalized, attr))
        xlsx_count = generated_wb[sheet_name].max_row - 1
        match = normalized_count == xlsx_count
        all_match = all_match and match
        print(f"{attr:<22} {normalized_count:>10} {xlsx_count:>8}  {'OK' if match else 'MISMATCH'}")
        assert normalized_count == xlsx_count, (
            f"{attr}: normalized reader returned {normalized_count} rows but "
            f"generated XLSX sheet {sheet_name!r} has {xlsx_count} — the API "
            f"and XLSX exporter are not reading the same canonical object."
        )
    contract_summary_in_xlsx = generated_wb["Contract Summary"].max_row - 1
    print(f"{'contract_summary':<22} {'1' if normalized.contract_summary else '0':>10} {contract_summary_in_xlsx:>8}  "
          f"{'OK' if bool(normalized.contract_summary) == (contract_summary_in_xlsx == 1) else 'MISMATCH'}")
    print(f"\nALL DATASETS CONSISTENT (normalized == xlsx): {all_match}")

    database.close()


if __name__ == "__main__":
    main()
