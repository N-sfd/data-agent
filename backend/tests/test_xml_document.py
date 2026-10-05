"""XML Document profile (xml_document@1): upload, safe parsing, record
detection, file-named columns, element-path provenance and exports."""

from __future__ import annotations

import io
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from app.database.session import SessionLocal
from app.main import app
from app.models.document import Document
from app.staging.preparation import prepare_staging
from app.xml_records.parse import XmlSourceError, parse_xml, tag_label

client = TestClient(app)

CLAUSE_XML = """<?xml version="1.0" encoding="UTF-8"?>
<OKCXMLIMPDFN xmlns="http://xmlns.oracle.com/apps/contracts/clauses" BatchName="FAR_PART_52">
  <Header><CreatedBy>Data Agent</CreatedBy></Header>
  <Clause>
    <Action>Sync</Action><DatePublished>2023-09-01</DatePublished><Number>52.201-1</Number>
    <Title>FAR 52.201-1 - Acquisition 360: Voluntary Survey</Title>
    <ProvisionYn>N</ProvisionYn><LockTextYn>Y</LockTextYn>
    <Text><p>Acquisition 360: Voluntary Survey (Sep 2023)</p><p>(a) All offerors are <b>encouraged</b> to provide feedback.</p></Text>
    <Attribute1>52.201-1</Attribute1>
  </Clause>
  <Clause>
    <Action>Sync</Action><DatePublished>2020-06-01</DatePublished><Number>52.202-1</Number>
    <Title>FAR 52.202-1 - Definitions</Title>
    <ProvisionYn>N</ProvisionYn><LockTextYn>Y</LockTextYn>
    <Text><p>Definitions (Jun 2020)</p></Text>
    <Attribute1>52.202-1</Attribute1>
  </Clause>
</OKCXMLIMPDFN>
"""


def test_tag_labels_follow_the_file():
    assert tag_label("ProvisionYn") == "Provision Yn"
    assert tag_label("Attribute1") == "Attribute 1"
    assert tag_label("DATE_PUBLISHED") == "Date Published"
    assert tag_label("XMLFileName") == "XML File Name"


def test_records_columns_paragraphs_and_details():
    parsed = parse_xml(CLAUSE_XML.encode())
    assert parsed.root_label == "OKCXMLIMPDFN"
    (group,) = parsed.groups
    assert group.label == "Clauses" and len(group.records) == 2
    assert [label for _, label in group.columns] == [
        "Action", "Date Published", "Number", "Title", "Provision Yn", "Lock Text Yn", "Text", "Attribute 1",
    ]
    text = group.records[0].values["Text"]
    # All <p> paragraphs stay together, one line each; inline markup is read inline.
    assert text.value == "Acquisition 360: Voluntary Survey (Sep 2023)\n(a) All offerors are encouraged to provide feedback."
    assert text.path == "/OKCXMLIMPDFN/Clause[1]/Text"
    assert {(d.label, d.value) for d in parsed.details} == {("Batch Name", "FAR_PART_52"), ("Created By", "Data Agent")}


def test_nested_structure_and_attributes_become_columns():
    raw = b"""<Orders>
      <Order id="A1"><Customer><Name>Acme</Name><City>Austin</City></Customer><Line sku="X">2</Line><Line sku="Y">3</Line></Order>
      <Order id="A2"><Customer><Name>Beta</Name><City>Boston</City></Customer><Line sku="Z">1</Line></Order>
    </Orders>"""
    (group,) = parse_xml(raw).groups
    labels = dict(group.columns)
    assert labels["@id"] == "ID" and labels["Customer/Name"] == "Customer Name"
    first = group.records[0].values
    assert first["Customer/City"].value == "Austin" and first["Line"].value == "2\n3"
    assert first["Line/@sku"].value == "X\nY" and labels["Line/@sku"] == "Line Sku"


def test_hostile_or_malformed_xml_is_refused():
    with pytest.raises(XmlSourceError):
        parse_xml(b'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x "boom">]><r>&x;</r>')
    with pytest.raises(XmlSourceError):
        parse_xml(b"<r><a></r>")


def _upload(content: bytes, name: str):
    return client.post("/api/documents/upload", files={"file": (name, content, "application/xml")})


def test_upload_rejects_malformed_xml():
    response = _upload(b"<r><a></r><!-- " + uuid4().hex.encode() + b" -->", "broken.xml")
    assert response.status_code in (400, 415, 422), response.text


def test_xml_staging_workbook_and_exports():
    from openpyxl import load_workbook

    content = CLAUSE_XML.replace("</OKCXMLIMPDFN>", f"<!-- {uuid4()} --></OKCXMLIMPDFN>").encode()
    upload = _upload(content, "FAR_Part_52_OKCXMLIMPDFN_PART1.xml")
    assert upload.status_code == 201, upload.text
    document_id = upload.json()["document_id"]
    pages = client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False})
    assert pages.status_code == 200, pages.text
    database = SessionLocal()
    try:
        preparation = prepare_staging(database, database.get(Document, document_id))
        assert preparation.record.profile_id == "xml_document"
    finally:
        database.close()

    workbook = client.get(f"/api/documents/{document_id}/staging-workbook").json()
    datasets = {d["dataset_id"]: d for d in workbook["datasets"]}
    # Unused record slots are not served.
    assert "xml_records_2" not in datasets
    clauses = datasets["xml_records_1"]
    assert clauses["display_name"] == "Clauses"
    assert [c["display_label"] for c in clauses["columns"]][:3] == ["Action", "Date Published", "Number"]
    cell = clauses["records"][1]["cells"][clauses["columns"][2]["canonical_field"]]
    assert cell["value"] == "52.202-1" and cell["review_status"] == "Verified"
    labels = [c["label"] for c in workbook["profile"]["export_capabilities"]]
    assert "Clauses CSV" in labels and not any(label.startswith("Records ") for label in labels)

    record = client.get(
        f"/api/documents/{document_id}/staging-workbook/datasets/xml_records_1/records/xml_records_1:1"
    ).json()
    title = next(c for c in record["cells"].values() if c["display_label"] == "Title")
    assert title["provenance"]["source_type"] == "xml"
    assert title["provenance"]["source_locator"]["dom_path"] == "/OKCXMLIMPDFN/Clause[1]/Title"

    csv_text = client.get(f"/api/documents/{document_id}/staging-workbook/datasets/xml_records_1.csv").content.decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith(
        "Action,Date Published,Number,Title,Provision Yn,Lock Text Yn,Text,Attribute 1,Source XML,Source Element"
    )

    xlsx = client.get(f"/api/documents/{document_id}/staging-workbook/export.xlsx")
    assert xlsx.status_code == 200
    book = load_workbook(io.BytesIO(xlsx.content))
    # The profile's workbook: the record group under its own name, details, source.
    assert book.sheetnames == ["CLAUSES", "DOCUMENT DETAILS", "SOURCE"]
    sheet = book["CLAUSES"]
    assert sheet["A1"].value == "XML Document — Clauses"
    assert "FAR_Part_52_OKCXMLIMPDFN_PART1.xml" in sheet["A2"].value
    headers = [c.value for c in sheet[4]]
    assert headers == ["Action", "Date Published", "Number", "Title", "Provision Yn", "Lock Text Yn", "Text", "Attribute 1", "Source XML"]
    assert sheet["H5"].value == "52.201-1" and sheet["I5"].value == "FAR_Part_52_OKCXMLIMPDFN_PART1.xml"
