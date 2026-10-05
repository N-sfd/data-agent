"""far_part_52@1 — FAR Part 52 regulation staging.

Synthetic FAR HTML (the acquisition.gov article/heading DOM) covers the
extraction rules; the reference comparison runs the real Part 52 HTML
against FAR_Part_52_Oracle_Transformation_Ready_v4.xlsx when both are
present locally (they are not required at runtime).
"""

from __future__ import annotations

import codecs
import io
import json
import re
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from app.database.session import SessionLocal
from app.far import views
from app.far.canonical import (
    ALTERNATE,
    CLAUSE_OR_PROVISION,
    NOTE_EMBEDDED,
    NOTE_RESERVED,
    NOTE_STRUCTURAL,
    RESERVED,
    SUBPART,
    build_canonical,
)
from app.contract_structure.builder import materialize_contract_structure
from app.far.dom import extract_far, recognize_far_part_52
from app.main import app
from app.models.document import Document
from app.staging import registry
from app.staging.preparation import prepare_staging
from app.staging.resolver import recognize_document_profile

client = TestClient(app)
ROOT = Path(__file__).resolve().parents[2]
REFERENCE_HTML = ROOT / "reference" / "regression" / "part_52.html"
REFERENCE_XLSX = ROOT / "reference" / "schemas" / "FAR_Part_52_Oracle_Transformation_Ready_v4.xlsx"


# --- synthetic FAR HTML ------------------------------------------------------------------


def _article(element_id: str, level: int, number: str, title: str, body: str = "", children: str = "") -> str:
    body_html = f'<div class="body conbody clause">{body}</div>' if body is not None else ""
    return (
        f'<article class="topic concept" id="{element_id}">'
        f'<h{level} class="title"><span class="ph autonumber">{number}</span> {title}</h{level}>'
        f"{body_html}{children}</article>"
    )


def _clause(
    number: str,
    title: str,
    date: str,
    extra: str = "",
    kind: str = "clause",
    prescribed: str = "9.999",
    after: str = "",
) -> str:
    element_id = "FAR_" + number.replace(".", "_").replace("-", "_")
    body = (
        f'<p class="p">As prescribed in <span class="ph ClauseProvisionXREF"><a class="xref" href="#x">{prescribed}</a>'
        f"(a)</span>, insert the following {kind}:</p>"
        f'<p class="p Ctr_SmCaps">{title} <span class="ph SmCaps">({date})</span></p>'
        f'<p class=" ListL1"><span class="ph autonumber">(a)</span> The Contractor shall comply.</p>'
        f"{extra}"
        f'<p class="p Endofclause">(End of {kind})</p>'
        f"{after}"
    )
    return _article(element_id, 4, number, f"{title}.", body)


def far_html(filler: int = 22) -> bytes:
    fillers = "".join(
        _clause(f"52.230-{n}", f"Filler Clause {n}", "Jan 2020") for n in range(1, filler + 1)
    )
    part_52_000 = _article(
        "FAR_52_000", 2, "52.000", "Scope of part.",
        '<p class="p">This part-</p><p class=" ListL1"><span class="ph autonumber">(a)</span> Gives instructions '
        "(e.g., fixed-price supply) for using provisions;</p>",
    )
    subpart_1 = _article(
        "FAR_Subpart_52_1", 2, "Subpart 52.1", "- Instructions for Using Provisions and Clauses", None,
        _article("FAR_52_100", 3, "52.100", "Scope of subpart.", '<p class="p">This subpart gives instructions (see 52.101).</p>'),
    )
    taxpayer = (
        '<article class="topic concept" id="FAR_52_204_3"><h4 class="title"><span class="ph autonumber">52.204-3</span>'
        ' Taxpayer Identification.</h4><div class="body conbody provision">'
        '<p class="p">As prescribed\n            in <a class="xref" href="#x">4.905</a>, insert the\n  following provision:</p>'
        '<p class="p Ctr_SmCaps">Taxpayer Identification <span class="ph SmCaps">(Oct 1998)</span></p>'
        '<p class=" ListL1"><span class="ph autonumber">(a)</span> <em>Definitions.</em></p>'
        '<p class=" ListL2"><span class="ph autonumber">(1)</span> Common parent means the entity; see '
        '<a class="xref" href="#FAR_52_204_6">52.204-6</a>.</p>'
        '<p class=" ListL3"><span class="ph autonumber">(i)</span> Nested item.</p>'
        '<p class=" ListL4"><span class="ph autonumber">(A)</span> Deeper item.</p>'
        '<table><tr><td>Item No.</td><td>Quantity</td></tr><tr><td>0001</td><td>10</td></tr></table>'
        '<p class="p Endofprovision">(End of provision)</p></div></article>'
    )
    uei = _clause("52.204-6", "Unique Entity Identifier", "NOV 2023", kind="provision", prescribed="4.607")
    narrative = _clause(
        "52.212-5", "Contract Terms and Conditions", "Jan 2025",
        extra='<p class=" ListL2"><span class="ph autonumber">(v)</span> Alternate IV (Jan 2025) of 52.219-9.</p>',
    )
    alternates = _clause(
        "52.215-1", "Instructions to Offerors-Competitive Acquisition", "Nov 2021", kind="provision", prescribed="15.209",
        after=(
            '<section class="section Alternate">'
            '<p class="p"><em class="ph i">Alternate I (Oct1997)</em>. As prescribed in <a class="xref" href="#x">15.209</a>'
            "(a)(1), substitute the following paragraph (f)(4) for paragraph (f)(4) of the basic provision:</p>"
            '<p class="p List1">(f)(4) The Government intends to evaluate proposals after discussions.</p>'
            '<p class="p"><em class="ph i">Alternate II (Oct 1997)</em>. As prescribed in 15.209(a)(2), add a paragraph (c)(9):</p>'
            '<p class="p List1">(c)(9) Offerors may submit proposals that depart from stated requirements (see 52.215-5).</p>'
            "</section>"
        ),
    )
    reserved_alt = _clause(
        "52.225-3", "Buy American-Free Trade Agreements", "Nov 2023",
        after='<p class="p">Alternate I [Reserved]</p>',
    )
    reserved = _article("FAR_52_203_1", 4, "52.203-1", "[Reserved]", "")
    subpart_2 = _article(
        "FAR_Subpart_52_2", 2, "Subpart 52.2", "- Text of Provisions and Clauses", None,
        _article("FAR_52_200", 3, "52.200", "Scope of subpart.", '<p class="p">This subpart sets forth the text.</p>')
        + reserved + taxpayer + uei + narrative + alternates + reserved_alt + fillers,
    )
    html = (
        '<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8"><title>FAR Part 52</title></head><body>'
        '<article class="nested0" id="FAR_Part_52"><h1 class="title"><span class="ph autonumber">Part 52</span>'
        " - Solicitation Provisions and Contract Clauses</h1>"
        '<div class="body"><p class=" ListL1"><a class="xref" href="#FAR_52_000">52.000 Scope of part.</a></p></div>'
        f"{part_52_000}{subpart_1}{subpart_2}</article>"
        f"<!-- {uuid4()} --></body></html>"
    )
    return html.encode("utf-8")


def contract_html_citing_far() -> bytes:
    """A contract that lists and quotes FAR clauses — must never resolve
    to FAR Part 52 (even when named part_52.html)."""

    listing = "".join(
        f"<tr><td>52.2{n:02d}-1</td><td>Clause Title {n}</td><td>NOV 2021</td></tr>" for n in range(10, 45)
    )
    full_text = "".join(
        f"<h3>52.2{n:02d}-4 Full Text Clause {n} (JAN 2020)</h3><p>The Contractor shall comply.</p>" for n in range(10, 35)
    )
    return (
        "<html><body><h1>CONTRACT W912DY-24-C-0001</h1><p>Contracting Officer: Jane Doe</p>"
        f"<h2>Section I - Contract Clauses</h2><table>{listing}</table>{full_text}"
        f"<!-- {uuid4()} --></body></html>"
    ).encode()


def _upload(content: bytes, name: str) -> str:
    response = client.post("/api/documents/upload", files={"file": (name, content, "text/html")})
    assert response.status_code == 201, response.text
    document_id = response.json()["document_id"]
    pages = client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False})
    assert pages.status_code == 200, pages.text
    return document_id


def _prepare(document_id: str) -> dict:
    database = SessionLocal()
    try:
        preparation = prepare_staging(database, database.get(Document, document_id))
        return {
            "profile_id": preparation.record.profile_id,
            "structure_built": preparation.structure_built,
            "timings_ms": preparation.timings_ms,
        }
    finally:
        database.close()


# --- extraction rules ---------------------------------------------------------------------


def _records():
    return {r.far_number: r for r in extract_far(far_html()).records}


def test_one_record_per_far_article_in_source_order():
    extraction = extract_far(far_html())
    numbers = [r.far_number for r in extraction.records]
    assert extraction.part_heading == "Part 52 - Solicitation Provisions and Contract Clauses"
    assert numbers[:6] == ["52.000", "Subpart 52.1", "52.100", "Subpart 52.2", "52.200", "52.203-1"]
    assert "Part 52" not in numbers
    assert len(numbers) == len(set(numbers))


def test_headings_dates_prescriptions_and_types():
    records = _records()
    taxpayer = records["52.204-3"]
    assert taxpayer.title == "Taxpayer Identification"
    assert taxpayer.official_heading == "Taxpayer Identification (Oct 1998)"
    assert taxpayer.revision_date == "Oct 1998"
    # Source line-wrapping inside the paragraph collapses; words are kept.
    assert taxpayer.prescription == "As prescribed in 4.905, insert the following provision:"
    assert taxpayer.prescription_reference == "4.905"
    assert taxpayer.clause_type == "Provision"
    # Upper-case month in the heading: date normalized, heading unchanged.
    uei = records["52.204-6"]
    assert uei.revision_date == "Nov 2023"
    assert uei.official_heading.endswith("(NOV 2023)")


def test_nested_paragraphs_and_tables_stay_inside_their_record():
    text = _records()["52.204-3"].basic_text
    for marker in ("(a) Definitions.", "(1) Common parent", "(i) Nested item.", "(A) Deeper item.", "Item No. | Quantity", "0001 | 10"):
        assert marker in text
    assert "52.204-3" not in [ref for ref in _records()["52.204-3"].embedded_references]


def test_text_is_never_rewritten():
    assert "(e.g., fixed-price supply)" in _records()["52.000"].basic_text


def test_structural_alternates_only_and_linked_to_their_basic_clause():
    records = _records()
    basic = records["52.215-1"]
    assert [a.code for a in basic.alternates] == ["Alternate I", "Alternate II"]
    assert [a.date for a in basic.alternates] == ["Oct 1997", "Oct 1997"]
    assert basic.alternates[0].instruction.startswith("As prescribed in 15.209")
    assert basic.alternates[0].prescription_reference == "15.209"
    assert "(f)(4) The Government intends" in basic.alternates[0].text
    # Not flattened into the basic text, and nothing lost.
    assert "Alternate I" not in basic.basic_text
    assert basic.basic_text + "\n" + "\n".join(a.text for a in basic.alternates) == basic.grouped_text
    # The narrative "(v) Alternate IV (Jan 2025) of 52.219-9" is not an alternate.
    narrative = records["52.212-5"]
    assert narrative.alternates == []
    assert narrative.narrative_alternate_mentions == 1
    # "Alternate I [Reserved]" is a structural (reserved) alternate.
    reserved = records["52.225-3"].alternates
    assert len(reserved) == 1 and reserved[0].reserved and reserved[0].date is None


def test_embedded_references_are_an_index_not_records():
    records = _records()
    assert records["52.204-3"].embedded_references == ["52.204-6"]
    assert records["52.100"].embedded_references == ["52.101"]
    assert "52.219-9" in records["52.212-5"].embedded_references
    assert "52.219-9" not in records and "52.101" not in records


def test_business_export_is_the_clause_library_layout_with_type_and_split_text():
    rows = build_canonical(extract_far(far_html()), "doc")
    by_key = {r.clause_key: views.business_row(r) for r in rows}
    assert list(views.BUSINESS_COLUMNS) == [
        "Action", "Date Published", "Number", "Title", "Display Name", "Type", "Intent", "Language",
        "Clause Type", "Status", "Description", "Provision Yn", "Global Yn", "Lock Text Yn",
        "Insert By Reference", "Text", "Provision Text", "Clause Text", "Start Date",
        "Attribute Category", "Attribute 1", "Source XML", "Source Reference",
    ]
    basic = by_key["FAR-52.204-3"]
    assert basic["Number"] == "52.204-3"
    assert basic["Title"].startswith("FAR 52.204-3 - ")
    assert basic["Type"] == "Provision" and basic["Provision Yn"] == "Y"
    assert basic["Text"] and "Taxpayer" in basic["Text"]
    # The prescription is the Description, not the start of the text.
    assert basic["Description"].startswith("As prescribed")
    assert not basic["Text"].startswith("As prescribed")
    assert basic["Provision Text"] == basic["Text"] and basic["Clause Text"] is None
    assert basic["Action"] == "Sync" and basic["Clause Type"] == "STANDARD" and basic["Attribute 1"] == "52.204-3"
    assert basic["Date Published"] and basic["Start Date"] == basic["Date Published"]
    assert basic["Source XML"] is None
    assert basic["Source Reference"] and "page" not in basic["Source Reference"].lower()
    reserved = by_key["FAR-52.203-1"]
    assert reserved["Type"] == "Reserved"
    # Records that are not loaded carry no library defaults.
    assert reserved["Action"] is None and reserved["Clause Type"] is None and reserved["Provision Yn"] is None
    alternate = by_key["FAR-52.215-1-ALT-I"]
    assert alternate["Type"] in ("Clause", "Provision") and alternate["Title"].endswith("Alternate I")
    assert alternate["Action"] is None
    # Part and subpart instructions (52.100 ...) are Sections with their text.
    section = by_key["FAR-52.100"]
    assert section["Type"] == "Section" and section["Text"]
    assert section["Provision Text"] is None and section["Clause Text"] is None
    assert all(views.business_row(r)["Source Reference"] for r in rows)
    assert views.published_date("Sep 2023") == "2023-09-01" and views.published_date("n/a") is None


def test_canonical_model_keys_types_and_load_eligibility():
    rows = build_canonical(extract_far(far_html()), "doc")
    by_key = {r.clause_key: r for r in rows}
    assert by_key["FAR-52.204-3"].content_type == CLAUSE_OR_PROVISION and by_key["FAR-52.204-3"].load_eligible
    assert by_key["FAR-Subpart_52.1"].content_type == SUBPART and not by_key["FAR-Subpart_52.1"].load_eligible
    assert by_key["FAR-Subpart_52.1"].transformation_notes == NOTE_STRUCTURAL
    assert by_key["FAR-52.203-1"].content_type == RESERVED and not by_key["FAR-52.203-1"].load_eligible
    assert by_key["FAR-52.203-1"].transformation_notes == NOTE_RESERVED
    assert NOTE_EMBEDDED in by_key["FAR-52.204-3"].transformation_notes
    alt = by_key["FAR-52.215-1-ALT-I"]
    assert alt.content_type == ALTERNATE and not alt.load_eligible
    assert alt.parent_clause_key == alt.basic_clause_key == "FAR-52.215-1"
    assert alt.version_date == "Oct 1997" and alt.alternate_code == "Alternate I"
    assert by_key["FAR-52.225-3-ALT-I"].content_type == RESERVED
    # Base records are their own basic clause; keys are deterministic.
    assert all(r.basic_clause_key == r.clause_key for r in rows if r.alternate_code is None)
    again = build_canonical(extract_far(far_html()), "doc")
    assert [r.clause_key for r in again] == [r.clause_key for r in rows]
    assert [r.source_order for r in rows] == list(range(1, len(rows) + 1))


def test_validation_flags_nothing_on_a_clean_source():
    rows = build_canonical(extract_far(far_html()), "doc")
    checks = {c["Check"]: c for c in views.validation_rows(rows)}
    assert all(c["Result"] == "PASS" for c in checks.values()), [c for c in checks.values() if c["Result"] != "PASS"]
    assert "1 narrative mention(s) ignored" in checks["Structural alternates only"]["Notes"]


# --- recognition / resolution -------------------------------------------------------------


def test_far_structure_is_recognized_and_contracts_are_not():
    assert recognize_far_part_52(far_html()).score >= 0.8
    assert recognize_far_part_52(contract_html_citing_far()).score < 0.8


def test_resolution_uses_structure_never_the_filename():
    far_id = _upload(far_html(), "regulation.html")
    contract_id = _upload(contract_html_citing_far(), "part_52.html")
    database = SessionLocal()
    try:
        far = recognize_document_profile(database, database.get(Document, far_id))
        assert far is not None and far.profile.key == "far_part_52@1" and far.confident
        # The FAR profile opts out of the contract pipeline (the job skips V3).
        assert far.profile.contract_pipeline is False
        assert recognize_document_profile(database, database.get(Document, contract_id)) is None
    finally:
        database.close()
    preparation = _prepare(contract_id)
    assert preparation["profile_id"] != "far_part_52"


def far_pdf() -> bytes:
    """A policy Part (not 52) printed from acquisition.gov: contents, then
    "Parent topic", then each heading above its own text."""

    import fitz

    contents, body = [], []
    for s in range(1, 4):
        subpart = f"Subpart 46.{s} - Subpart {s} Title"
        contents.append(subpart)
        body.append(subpart)
        for n in range(1, 5):
            heading = f"46.{s}0{n} Section {s}{n} title."
            contents.append(heading)
            body += [heading, f"(a) Text of section {s}{n} cites 52.246-2 and 46.10{n}."]
    paragraphs = ["Part 46 - Quality Assurance", *contents, "Parent topic: Federal Acquisition Regulation", *body]
    document = fitz.open()
    page, y = document.new_page(), 40
    for text in paragraphs:
        if y > 760:
            page, y = document.new_page(), 40
        page.insert_text((50, y), text, fontsize=9)
        y += 22
    return document.tobytes()


def test_far_part_pdf_resolves_to_the_far_profile_with_records():
    response = client.post("/api/documents/upload", files={"file": ("part-46.pdf", far_pdf(), "application/pdf")})
    assert response.status_code == 201, response.text
    document_id = response.json()["document_id"]
    assert client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False}).status_code == 200
    assert _prepare(document_id)["profile_id"] == "far_part_52"
    workbook = _workbook(document_id)
    records = next(d for d in workbook["datasets"] if d["dataset_id"] == "far_records")["records"]
    numbers = [r["cells"]["far.record.far_number"]["value"] for r in records]
    assert numbers[:3] == ["Subpart 46.1", "46.101", "46.102"] and len(numbers) == 15
    qa = {
        r["cells"]["qa.check"]["value"]: r["cells"]["qa.result"]["value"]
        for r in next(d for d in workbook["datasets"] if d["dataset_id"] == "qa_review")["records"]
    }
    assert qa["Stable unique clause keys"] == "PASS"


def test_existing_profiles_keep_their_capabilities():
    for key in ("contract_v3", "invoice", "generic_business_document"):
        profile = registry.latest(key)
        assert profile.document_recognizer is None
        # Contracts materialize their source-adaptive structure once; the
        # other pre-FAR profiles still build nothing up front.
        if key == "contract_v3":
            assert profile.materializer is materialize_contract_structure
        else:
            assert profile.materializer is None
        assert profile.contract_pipeline is True
        if key != "contract_v3":
            assert all(not d.grid_fields for d in profile.datasets)


# --- staging workbook, record details and exports -----------------------------------------


@pytest.fixture(scope="module")
def far_document() -> str:
    document_id = _upload(far_html(), "far_part_52_sample.html")
    preparation = _prepare(document_id)
    assert preparation["profile_id"] == "far_part_52"
    assert preparation["structure_built"] is False
    assert "profile_materialize_ms" in preparation["timings_ms"]
    return document_id


def _workbook(document_id: str) -> dict:
    response = client.get(f"/api/documents/{document_id}/staging-workbook")
    assert response.status_code == 200, response.text
    return response.json()


def test_workbook_tabs_and_states(far_document):
    workbook = _workbook(far_document)
    profile = workbook["profile"]
    assert profile["profile_id"] == "far_part_52"
    # Three business-facing tabs; QA sits with the source.
    assert [(v["label"], v["dataset_ids"]) for v in profile["views"]] == [
        ("FAR Clauses & Provisions", ["far_records"]),
        ("Oracle Output", ["far_oracle_output"]),
        ("Source", ["far_source", "qa_review"]),
    ]
    assert [d["dataset_id"] for d in workbook["datasets"]] == ["far_records", "far_oracle_output", "far_source", "qa_review"]
    assert workbook["qa_summary"]["needs_review"] == 0 and workbook["qa_summary"]["verified"] > 0
    records = next(d for d in workbook["datasets"] if d["dataset_id"] == "far_records")
    # The Part 52 columns in order; the FAR_REGULATION fields other Parts
    # use (Part / Subpart titles, Part 53 forms) are blank here.
    labels = [c["display_label"] for c in records["columns"]]
    assert [label for label in labels if label in views.FAR_RECORD_COLUMNS] == list(views.FAR_RECORD_COLUMNS)
    assert records["display_name"] == "FAR Clauses & Provisions"
    assert records["compact"] is True
    # Grid rows carry values and states only — no provenance.
    row = records["records"][0]
    assert all(cell["provenance"] is None for cell in row["cells"].values())
    oracle = next(d for d in workbook["datasets"] if d["dataset_id"] == "far_oracle_output")
    assert oracle["role"] == "transform"
    assert [c["display_label"] for c in oracle["columns"]] == list(views.ORACLE_OUTPUT_COLUMNS)
    # The Oracle shape is a view of the same records, not more records.
    assert workbook["outcome"]["record_count"] == len(records["records"])


def _far_rows(workbook: dict) -> dict[str, dict]:
    records = next(d for d in workbook["datasets"] if d["dataset_id"] == "far_records")
    rows = {}
    for record in records["records"]:
        values = {cell["display_label"]: cell["value"] for cell in record["cells"].values()}
        rows[" ".join(v for v in (values["FAR Number"], values["Alternate"]) if v)] = values
    return rows


def test_far_clauses_and_provisions_follow_far_semantics(far_document):
    rows = _far_rows(_workbook(far_document))
    taxpayer = rows["52.204-3"]
    assert taxpayer["Title"] == "Taxpayer Identification"
    assert taxpayer["Record Type"] == "Provision" and taxpayer["Record Status"] == "Active"
    assert (taxpayer["FAR Part"], taxpayer["FAR Subpart"], taxpayer["FAR Section"]) == ("Part 52", "Subpart 52.2", "52.204")
    assert taxpayer["Revision Date"] == "Oct 1998"
    assert taxpayer["Description"] == "Taxpayer Identification (Oct 1998)"
    assert taxpayer["Prescription / Usage"].startswith("As prescribed")
    assert taxpayer["Prescription Reference"] == "4.905"
    assert taxpayer["Cross References"] == "52.204-6"
    # A provision's body only in Provision Text, without its prescription.
    assert taxpayer["Provision Text"].startswith("Taxpayer Identification (Oct 1998)")
    assert taxpayer["Clause Text"] is None
    assert taxpayer["Source Reference"].startswith("far_part_52_sample.html › Subpart 52.2")
    assert taxpayer["Source Reference"].endswith("#FAR_52_204_3")

    clause = rows["52.212-5"]
    assert clause["Record Type"] == "Clause" and clause["Provision Text"] is None
    assert clause["Prescription Reference"] == "9.999(a)"
    assert "Contract Terms and Conditions" in clause["Clause Text"]

    # Alternates keep their FAR Number and their parent's type.
    alt = rows["52.215-1 Alternate I"]
    assert alt["Record Type"] == "Provision" and alt["Title"] == rows["52.215-1"]["Title"]
    assert alt["Prescription Reference"] == "15.209(a)(1)"
    assert alt["Provision Text"] and alt["Clause Text"] is None
    # Its heading and instruction have their own columns; the body follows them.
    assert alt["Description"].startswith("Alternate I")
    assert alt["Provision Text"].startswith("(f)(4) The Government intends")
    assert rows["52.215-1 Alternate II"]["Prescription Reference"] == "15.209(a)(2)"

    # Structural records: no body columns, no invented statuses.
    # A Part 52 section is classified by its regulatory function.
    assert rows["52.000"]["Record Type"] == "Scope"
    assert rows["52.000"]["Provision Text"] is None and rows["52.000"]["Clause Text"] is None
    assert rows["52.203-1"]["Record Status"] == "Reserved"
    assert rows["Subpart 52.1"]["Record Type"] == "Subpart"
    assert all(r["Record Status"] in ("Active", "Reserved") for r in rows.values())


def test_record_details_carry_text_and_dom_provenance(far_document):
    response = client.get(
        f"/api/documents/{far_document}/staging-workbook/datasets/far_records/records/far_records:FAR-52.204-3"
    )
    assert response.status_code == 200, response.text
    cells = response.json()["cells"]
    text = cells["far.record.provision_text"]["value"]
    assert "(A) Deeper item." in text
    heading = cells["far.record.description"]
    assert heading["review_status"] == "Verified"
    locator = heading["provenance"]["source_locator"]
    assert locator["element_id"] == "FAR_52_204_3" and locator["dom_path"].startswith("/html/body")
    assert locator["section_path"][-1] == "52.204-3 Taxpayer Identification."
    # HTML has no pages and no coordinates — none are invented.
    assert heading["provenance"]["source_page"] is None and heading["provenance"]["source_bbox"] is None
    missing = client.get(f"/api/documents/{far_document}/staging-workbook/datasets/far_records/records/nope")
    assert missing.status_code == 404


def test_exports_workbook_csv_and_json(far_document):
    from openpyxl import load_workbook

    base = f"/api/documents/{far_document}/staging-workbook/exports"
    xlsx = client.get(f"{base}/far_workbook.xlsx")
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content), read_only=True)
    # LONG TEXT only exists when a value is longer than an Excel cell.
    assert wb.sheetnames == ["FAR CLAUSES & PROVISIONS", "ORACLE OUTPUT", "SOURCE"]
    far = list(wb["FAR CLAUSES & PROVISIONS"].iter_rows(values_only=True))
    assert far[0][0] == "FAR Part 52 — Clauses & Provisions"
    assert list(far[3]) == list(views.FAR_RECORD_COLUMNS)
    oracle = list(wb["ORACLE OUTPUT"].iter_rows(values_only=True))
    assert list(oracle[3]) == list(views.ORACLE_OUTPUT_COLUMNS)
    source = {r[0]: r[1] for r in list(wb["SOURCE"].iter_rows(values_only=True))[4:]}
    assert source["Source Document"] == "far_part_52_sample.html"
    assert source["Alternates (each follows its basic record)"] >= 2
    assert source["Records (incl. alternates)"] == len(far) - 4

    technical = load_workbook(io.BytesIO(client.get(f"{base}/far_transformation.xlsx").content), read_only=True)
    assert technical.sheetnames[0] == "01_FAR_STAGING" and "08_BUSINESS_EXPORT" in technical.sheetnames

    csv = client.get(f"{base}/far_clauses_provisions.csv")
    # UTF-8 with a byte-order mark, so Excel does not show "›" as "â€º".
    assert csv.status_code == 200 and csv.content.startswith(codecs.BOM_UTF8)
    assert csv.headers["content-type"] == "text/csv; charset=utf-8"
    assert csv.content.decode("utf-8-sig").startswith("FAR Number,Title,Record Type,FAR Part")
    assert client.get(f"{base}/oracle_output.csv").content.decode("utf-8-sig").startswith("Action,Date Published,Number")
    assert client.get(f"{base}/06_CANONICAL_MODEL.csv").content.decode("utf-8-sig").startswith("Source Sequence ID,Clause Key")
    payload = json.loads(client.get(f"{base}/far_canonical.json").text)
    assert payload["profile"] == "far_part_52@1"
    assert {r["Clause Key"] for r in payload["canonical_model"]} >= {"FAR-52.204-3", "FAR-52.215-1-ALT-I"}
    assert len(payload["far_clauses_provisions"]) == len(payload["canonical_model"])
    assert client.get(f"{base}/unknown.bin").status_code == 404


def test_long_text_sheet_appears_only_when_a_cell_overflows():
    from openpyxl import load_workbook

    from app.far.exports import build_far_workbook
    from app.models.document_far_record import DocumentFarRecord

    body = "As prescribed in 1.101, insert the following clause:\nLong Clause (Jan 2020)\n" + "x" * 40_000
    row = DocumentFarRecord(
        clause_key="FAR-52.299-1", basic_clause_key="FAR-52.299-1", source_sequence_id="1", source_order=1,
        content_type="CLAUSE_OR_PROVISION", load_eligible=True, far_number="52.299-1", title="Long Clause",
        clause_type="Clause", prescription="As prescribed in 1.101, insert the following clause:",
        source_text=body, embedded_references=[], provenance_json={}, issues=[], extractor_version=1,
    )
    wb = load_workbook(io.BytesIO(build_far_workbook([row], {}, "long.html")), read_only=True)
    assert wb.sheetnames[-1] == "LONG TEXT"
    cell = list(wb["FAR CLAUSES & PROVISIONS"].iter_rows(values_only=True))[4][14]
    assert "continued in LONG TEXT" in cell
    parts = [p for p in list(wb["LONG TEXT"].iter_rows(values_only=True))[4:] if p[0] == "FAR CLAUSES & PROVISIONS"]
    assert [p[3] for p in parts] == ["1 of 2", "2 of 2"] and parts[0][1:3] == ("52.299-1", "Clause Text")


def test_other_profiles_have_no_profile_exports():
    contract_id = _upload(contract_html_citing_far(), "contract.html")
    _prepare(contract_id)
    assert client.get(f"/api/documents/{contract_id}/staging-workbook/exports/far_workbook.xlsx").status_code == 404


# --- reference comparison (local only) ----------------------------------------------------

REPRESENTATIVE = (
    "52.000", "Subpart 52.1", "52.100", "52.101", "52.204-3", "52.204-6",
    "52.207-4", "52.209-1", "52.211-3", "52.211-8", "52.211-9", "52.212-3",
)


def _ns(text) -> str:
    return re.sub(r"\s+", "", str(text or ""))


@pytest.mark.skipif(
    not (REFERENCE_HTML.exists() and REFERENCE_XLSX.exists()),
    reason="FAR reference HTML/workbook not available locally",
)
def test_reference_comparison_with_transformation_workbook():
    from openpyxl import load_workbook

    rows = build_canonical(extract_far(REFERENCE_HTML.read_bytes()), "reference")
    base = views.base_rows(rows)
    by_number = {r.far_number: r for r in base}
    wb = load_workbook(REFERENCE_XLSX, read_only=True)
    structured = [r for r in wb["04_STRUCTURED_FAR"].iter_rows(min_row=5, values_only=True) if r[0]]
    canonical = [r for r in wb["06_CANONICAL_MODEL"].iter_rows(min_row=5, values_only=True) if r[0]]

    # Records, order and identity.
    assert [r.far_number for r in base] == [r[1] for r in structured]
    assert len(base) == len(canonical) == 760
    assert [r.clause_key for r in base] == [r[1] for r in canonical]
    assert [r.source_sequence_id for r in base] == [r[0] for r in canonical]

    count = lambda predicate, items: sum(1 for item in items if predicate(item))
    # Content types and load eligibility match exactly.
    for content_type in (CLAUSE_OR_PROVISION, RESERVED, SUBPART):
        assert count(lambda r: r.content_type == content_type, base) == count(lambda r: r[4] == content_type, canonical)
    assert count(lambda r: r.load_eligible, base) == count(lambda r: r[12] == "YES", canonical) == 621
    assert count(lambda r: r.basic_clause_key == r.clause_key, base) == 760
    assert count(lambda r: r.parent_clause_key, base) == 0
    assert count(lambda r: r.source_text, base) == count(lambda r: r[13], canonical) == 621
    assert count(lambda r: r.embedded_references, base) == count(lambda r: r[14], canonical) == 84
    assert count(lambda r: views.structured_row(r)["Official Heading"], base) == count(lambda r: r[4], structured) == 613
    for note in (NOTE_RESERVED, NOTE_STRUCTURAL, NOTE_EMBEDDED):
        assert count(lambda r: note in (r.transformation_notes or ""), base) == count(lambda r: r[15] == note, canonical)

    # Every value the reference has, we have identically (whitespace aside);
    # where the reference has none, ours comes from the source heading.
    for ref in structured:
        ours = by_number[ref[1]]
        if ref[5]:
            assert ours.version_date.replace("\xa0", " ") == ref[5].replace("\xa0", " ")
        if ref[6]:
            assert _ns(ours.prescription).replace("e.g.,", "").replace("e.g.", "").startswith(_ns(ref[6]))
        if ref[4] and ours.content_type != SUBPART:
            assert _ns(ours.official_heading) == _ns(ref[4]).replace("�", "—")
        if ref[2] and ref[2] != "Subpart":
            # The reference cuts titles at a source line break; ours is whole.
            assert _ns(ours.title).startswith(_ns(ref[2]).replace("�", "—"))
    for ref in canonical:
        if ref[10]:
            assert by_number[ref[3]].prescription_reference == ref[10]
    # Reference text equals ours after the reference's own cleanups (it
    # removes "e.g.", joins table cells, and continues long text in 99).
    long_or_reserved_alt = 0
    for ref in structured:
        ours = by_number[ref[1]]
        expected = _ns(ref[11])
        actual = _ns(ours.source_text).replace("|", "").replace("e.g.,", "").replace("e.g.", "")
        if expected != actual:
            assert actual.startswith(expected.split("[Fulltextcontinues")[0])
            long_or_reserved_alt += 1
    assert long_or_reserved_alt <= 6

    # Structural alternates: the reference has none; the source has 204
    # (5 reserved placeholders) — each linked, none from narrative mentions.
    alternates = [r for r in rows if r.alternate_code]
    assert len(alternates) == 204
    assert all(a.basic_clause_key in {b.clause_key for b in base} for a in alternates)

    # Representative records.
    for number in REPRESENTATIVE:
        assert number in by_number
    assert by_number["52.204-3"].prescription_reference == "4.905"  # reference missed it (line-wrapped)
    assert by_number["52.211-3"].version_date == "June 1988"  # reference missed four-letter months
    assert by_number["52.204-6"].prescription_reference == "4.607"
    assert by_number["52.212-3"].version_date == "Oct 2025"
    assert len(by_number["52.212-3"].source_text) > 30000  # whole, not truncated
    assert by_number["Subpart 52.1"].content_type == SUBPART


# --- derived values are Verified only when the derivation is unambiguous -------------------


def test_derivation_checks_flag_ambiguous_sources():
    from app.models.document_far_record import DocumentFarRecord
    from app.staging.profiles.far_part_52 import (
        _clause_type_check,
        _date_check,
        _prescription_reference_check,
    )

    # One prescribing section → unambiguous; two → Needs Review.
    assert _prescription_reference_check("4.905", "As prescribed in 4.905, insert the following provision:")[0].passed
    two = _prescription_reference_check("32.205", "As prescribed in 32.205 (b) and 32.206 , insert the following clause:")
    assert not two[0].passed and "more than one prescribing section" in two[0].message

    # Prescription and end marker must agree on clause vs provision.
    agree = DocumentFarRecord(
        clause_type="Provision",
        prescription="As prescribed in 4.905, insert the following provision:",
        source_text="As prescribed ...\n(a) Text.\n(End of provision)",
    )
    assert _clause_type_check(agree)[0].passed
    conflict = DocumentFarRecord(
        clause_type="Clause",
        prescription="As prescribed in 9.206-2, insert the following clause:",
        source_text="As prescribed ...\n(a) Text.\n(End of provision)",
    )
    assert not _clause_type_check(conflict)[0].passed
    either = DocumentFarRecord(
        clause_type="Provision",
        prescription="As prescribed in 28.101-2 , insert a provision or clause substantially as follows:",
        source_text="",
    )
    assert not _clause_type_check(either)[0].passed

    # A heading date: exactly one (Month YYYY), matching.
    assert _date_check("Nov 2023", "Buy American (NOV 2023)")[0].passed
    assert not _date_check("Oct 2016", "Unique Entity Identifier (Nov 2021)")[0].passed


@pytest.mark.skipif(not REFERENCE_HTML.exists(), reason="FAR reference HTML not available locally")
def test_reference_source_ambiguities_are_needs_review_not_verified():
    from app.staging.profiles.far_part_52 import _clause_type_check, _prescription_reference_check

    rows = views.base_rows(build_canonical(extract_far(REFERENCE_HTML.read_bytes()), "reference"))
    by_number = {r.far_number: r for r in rows}
    multi = by_number["52.232-31"]
    assert not _prescription_reference_check(multi.prescription_reference, multi.prescription)[0].passed
    assert not _clause_type_check(by_number["52.209-1"])[0].passed  # "insert the following clause" / "(End of provision)"
    assert _clause_type_check(by_number["52.204-6"])[0].passed


def test_grid_carries_a_text_preview_and_names_the_truncated_fields():
    from app.staging.models import StagingCell, StagingDataset, StagingRecord
    from app.staging.service import GRID_TEXT_PREVIEW, compact_for_grid

    long_text = "word " * 400
    record = StagingRecord(
        record_id="r1",
        cells={
            "far.business.text": StagingCell(canonical_field="far.business.text", display_label="Text", value=long_text),
            "far.business.number": StagingCell(canonical_field="far.business.number", display_label="Number", value="52.101"),
        },
    )
    dataset = StagingDataset(
        dataset_id="far_business", display_name="Business Export", cardinality="repeating", columns=[],
        records=[record], grid_fields=["far.business.text", "far.business.number"],
    )

    class _Workbook:
        datasets = [dataset]

    compact_for_grid(_Workbook())
    cell = dataset.records[0].cells["far.business.text"]
    assert len(cell.value) <= GRID_TEXT_PREVIEW + 1 and cell.value.endswith("…")
    assert dataset.records[0].truncated_fields == ["far.business.text"]
    assert dataset.records[0].cells["far.business.number"].value == "52.101"
    assert "truncated_fields" in dataset.records[0].model_dump()
    plain = StagingRecord(record_id="r2", cells={})
    assert "truncated_fields" not in plain.model_dump()


def test_full_text_grid_keeps_complete_values():
    from app.staging.models import StagingCell, StagingDataset, StagingRecord
    from app.staging.service import compact_for_grid

    long_text = "word " * 400
    record = StagingRecord(
        record_id="r1",
        cells={"far.record.clause_text": StagingCell(canonical_field="far.record.clause_text", display_label="Clause Text", value=long_text)},
    )
    dataset = StagingDataset(
        dataset_id="far_records", display_name="FAR Clauses & Provisions", cardinality="repeating", columns=[],
        records=[record], grid_fields=["far.record.clause_text"], full_text=True,
    )

    class _Workbook:
        datasets = [dataset]

    compact_for_grid(_Workbook())
    assert dataset.records[0].cells["far.record.clause_text"].value == long_text
    assert dataset.records[0].truncated_fields == []


def test_far_records_tab_serves_complete_texts(far_document):
    records = next(d for d in _workbook(far_document)["datasets"] if d["dataset_id"] == "far_records")
    assert records["full_text"] is True
    taxpayer = next(r for r in records["records"] if r["record_id"] == "far_records:FAR-52.204-3")
    assert "truncated_fields" not in taxpayer
    assert taxpayer["cells"]["far.record.provision_text"]["value"].rstrip().endswith("(End of provision)")
