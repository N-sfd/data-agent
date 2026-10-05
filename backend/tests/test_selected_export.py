"""Selected export (POST …/staging-workbook/datasets/{id}/export-selected):
the cells a reviewer picked, as Excel / CSV / JSON, read from the persisted
staging workbook — table shape kept, identifier carried, only business
fields, and every record checked against the document's own dataset."""

from __future__ import annotations

import codecs
import csv
import io
import json
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.models.document import Document
from app.staging import registry
from app.staging.engine import assemble_workbook
from app.staging.models import ProcessingMetadata
from app.staging.profile import AdapterResult, RawRecord
from app.staging.selection_export import SelectionError, SelectionManifest, build_selected_export, select

client = TestClient(app)

F = "far.record."


def _upload_far_pdf() -> str:
    from tests.test_far_part_52 import _prepare, far_pdf

    content = far_pdf() + b"\n%" + uuid4().hex.encode()
    response = client.post("/api/documents/upload", files={"file": ("Part-46---Quality-Assurance.pdf", content, "application/pdf")})
    assert response.status_code == 201, response.text
    document_id = response.json()["document_id"]
    assert client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False}).status_code == 200
    assert _prepare(document_id)["profile_id"] == "far_part_52"
    return document_id


@pytest.fixture(scope="module")
def far() -> dict:
    document_id = _upload_far_pdf()
    workbook = client.get(f"/api/documents/{document_id}/staging-workbook").json()
    records = next(d for d in workbook["datasets"] if d["dataset_id"] == "far_records")["records"]
    by_number = {r["cells"][F + "far_number"]["value"]: r for r in records}
    return {"id": document_id, "records": records, "by_number": by_number}


def _export(document_id: str, dataset_id: str, fmt: str, selection: list[dict]):
    return client.post(
        f"/api/documents/{document_id}/staging-workbook/datasets/{dataset_id}/export-selected",
        json={"format": fmt, "selection": selection},
    )


def _csv(response) -> list[list[str]]:
    assert response.status_code == 200, response.text
    assert response.content.startswith(codecs.BOM_UTF8)
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def _rid(far: dict, number: str) -> str:
    return far["by_number"][number]["record_id"]


def _value(far: dict, number: str, key: str):
    return far["by_number"][number]["cells"][F + key]["value"]


# --- selection shapes -------------------------------------------------------------------


def test_one_cell_keeps_the_identifier(far):
    rows = _csv(_export(far["id"], "far_records", "csv", [{"record_id": _rid(far, "46.202"), "fields": [F + "section_text"]}]))
    assert rows == [["FAR Number", "Description / Regulatory Text"], ["46.202", _value(far, "46.202", "section_text")]]


def test_sparse_cells_across_records_leave_unselected_cells_empty(far):
    selection = [
        {"record_id": _rid(far, "46.101"), "fields": [F + "title"]},
        {"record_id": _rid(far, "46.102"), "fields": [F + "section_text"]},
    ]
    rows = _csv(_export(far["id"], "far_records", "csv", selection))
    assert rows[0] == ["FAR Number", "Title", "Description / Regulatory Text"]
    assert rows[1] == ["46.101", _value(far, "46.101", "title"), ""]
    assert rows[2] == ["46.102", "", _value(far, "46.102", "section_text")]


def test_rows_keep_source_order_and_columns_keep_dataset_order(far):
    # Requested out of order: records and fields come back in source / dataset order.
    fields = [F + "section_text", F + "title", F + "far_number"]
    selection = [{"record_id": _rid(far, n), "fields": fields} for n in ("46.103", "46.101", "46.102")]
    rows = _csv(_export(far["id"], "far_records", "csv", selection))
    assert rows[0] == ["FAR Number", "Title", "Description / Regulatory Text"]
    assert [r[0] for r in rows[1:]] == ["46.101", "46.102", "46.103"]
    assert rows[2] == ["46.102", _value(far, "46.102", "title"), _value(far, "46.102", "section_text")]


def test_one_row_exports_every_selected_field(far):
    fields = [c for c in far["by_number"]["46.201"]["cells"] if c.startswith(F)]
    payload = json.loads(_export(far["id"], "far_records", "json", [{"record_id": _rid(far, "46.201"), "fields": fields}]).content)
    assert len(payload["records"]) == 1
    record = payload["records"][0]
    from app.far.regulation import RECORD_TYPES

    assert record["far_number"] == "46.201" and record["record_type"] in RECORD_TYPES
    assert record["section_text"] == _value(far, "46.201", "section_text")


def test_one_column_over_many_records(far):
    numbers = [n for n in far["by_number"] if not n.startswith("Subpart")]
    selection = [{"record_id": _rid(far, n), "fields": [F + "title"]} for n in numbers]
    rows = _csv(_export(far["id"], "far_records", "csv", selection))
    assert rows[0] == ["FAR Number", "Title"]
    assert len(rows) - 1 == len(numbers)


def test_json_uses_semantic_field_names(far):
    response = _export(far["id"], "far_records", "json", [{"record_id": _rid(far, "46.102"), "fields": [F + "title", F + "section_text"]}])
    assert response.status_code == 200 and response.headers["content-type"].startswith("application/json")
    payload = json.loads(response.content)
    assert payload["document"].startswith("Part-46---Quality-Assurance") and payload["dataset"] == "FAR Data"
    assert payload["records"] == [
        {"far_number": "46.102", "title": _value(far, "46.102", "title"), "section_text": _value(far, "46.102", "section_text")}
    ]
    assert 'filename="Part-46---Quality-Assurance_far_records_selected.json"' in response.headers["content-disposition"]


def test_excel_keeps_the_professional_layout(far):
    selection = [{"record_id": _rid(far, n), "fields": [F + "far_number", F + "title", F + "section_text"]} for n in ("46.101", "46.102")]
    response = _export(far["id"], "far_records", "xlsx", selection)
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet["A1"].value == "FAR Part 46 — Selected Records"
    assert sheet["A2"].value.startswith("Extracted from Part-46---Quality-Assurance")
    assert [c.value for c in sheet[4]] == ["FAR Number", "Title", "Description / Regulatory Text"]
    assert [c.value for c in sheet[5]] == ["46.101", _value(far, "46.101", "title"), _value(far, "46.101", "section_text")]
    assert sheet.freeze_panes == "A5" and sheet.auto_filter.ref == "A4:C6"
    assert sheet["C5"].alignment.wrap_text


def test_values_are_the_persisted_workbook_values(far):
    """Exports read the persisted records: the workbook's own values,
    never anything the request carried."""

    record = far["by_number"]["46.103"]
    fields = [F + "title", F + "record_type", F + "far_subpart", F + "section_text"]
    payload = json.loads(_export(far["id"], "far_records", "json", [{"record_id": record["record_id"], "fields": fields}]).content)
    exported = payload["records"][0]
    for field in fields:
        assert exported[field.removeprefix(F)] == record["cells"][field]["value"]


def test_other_datasets_of_the_profile_export_too(far):
    oracle = next(
        d for d in client.get(f"/api/documents/{far['id']}/staging-workbook").json()["datasets"] if d["dataset_id"] == "far_oracle_output"
    )
    record = oracle["records"][3]
    rows = _csv(_export(far["id"], "far_oracle_output", "csv", [{"record_id": record["record_id"], "fields": ["far.oracle.text"]}]))
    assert rows[0] == ["Number", "Text"]
    assert rows[1] == [record["cells"]["far.oracle.number"]["value"], record["cells"]["far.oracle.text"]["value"] or ""]


# --- authorization / validation ---------------------------------------------------------


def test_unknown_record_ids_are_rejected(far):
    response = _export(far["id"], "far_records", "csv", [{"record_id": "far_records:FAR-99.999", "fields": [F + "title"]}])
    assert response.status_code == 422 and "not in this document" in response.json()["detail"]


def test_records_of_another_dataset_are_rejected(far):
    # A record id is checked against the dataset in the URL (and so the
    # document): FAR Clauses & Provisions ids are not QA Review records.
    response = _export(far["id"], "qa_review", "csv", [{"record_id": _rid(far, "46.101"), "fields": ["qa.check"]}])
    assert response.status_code == 422


def test_another_workspace_cannot_export_the_document():
    import secrets

    from tests.test_document_isolation import enforced
    from tests.test_far_part_52 import far_pdf

    a, b = {"X-Workspace-Token": secrets.token_hex(32)}, {"X-Workspace-Token": secrets.token_hex(32)}
    with enforced():
        content = far_pdf() + b"\n%" + uuid4().hex.encode()
        upload = client.post("/api/documents/upload", files={"file": ("p46.pdf", content, "application/pdf")}, headers=a)
        document_id = upload.json()["document_id"]
        client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False}, headers=a)
        records = next(
            d for d in client.get(f"/api/documents/{document_id}/staging-workbook", headers=a).json()["datasets"]
            if d["dataset_id"] == "far_records"
        )["records"]
        body = {"format": "csv", "selection": [{"record_id": records[1]["record_id"], "fields": [F + "title"]}]}
        url = f"/api/documents/{document_id}/staging-workbook/datasets/far_records/export-selected"
        assert client.post(url, json=body, headers=a).status_code == 200
        assert client.post(url, json=body, headers=b).status_code == 404


def test_unknown_or_technical_fields_are_rejected(far):
    for field in ("far.record.bbox", "far.record.confidence", "record_id", "far.oracle.text"):
        response = _export(far["id"], "far_records", "csv", [{"record_id": _rid(far, "46.101"), "fields": [field]}])
        assert response.status_code == 422, field


def test_unknown_dataset_document_and_empty_selection(far):
    assert _export(far["id"], "nope", "csv", [{"record_id": "x", "fields": ["y"]}]).status_code == 404
    assert _export(str(uuid4()), "far_records", "csv", [{"record_id": "x", "fields": ["y"]}]).status_code == 404
    assert _export(far["id"], "far_records", "csv", []).status_code == 422
    assert _export(far["id"], "far_records", "pdf", [{"record_id": _rid(far, "46.101"), "fields": [F + "title"]}]).status_code == 422


def test_export_all_is_unchanged(far):
    response = client.get(f"/api/documents/{far['id']}/staging-workbook/exports/far_clauses_provisions.csv")
    assert response.status_code == 200
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert len(rows) - 1 == len(far["records"])


# --- long text and other profiles (the same component for every table) -------------------


def _workbook_for(profile_id: str, dataset_id: str, records: list[RawRecord]):
    profile = registry.get_profile(profile_id)
    document = Document(id="doc-sel", original_filename="Source Document.pdf", stored_filename="doc-sel.pdf")
    return assemble_workbook(
        profile=profile,
        document=document,
        adapter_result=AdapterResult(records={dataset_id: records}),
        metadata=ProcessingMetadata(),
    )


def test_long_clause_text_stays_complete():
    long_text = "(a) " + "The Contractor shall comply. " * 2_000  # ~58,000 characters
    workbook = _workbook_for(
        "far_part_52",
        "far_records",
        [RawRecord("r1", {"far_number": "52.212-4", "title": "Contract Terms", "clause_text": long_text})],
    )
    selection = [{"record_id": "r1", "fields": [F + "far_number", F + "clause_text"]}]
    csv_content, _, _ = build_selected_export(workbook, "far_records", SelectionManifest(format="csv", selection=selection))
    assert list(csv.reader(io.StringIO(csv_content)))[1][1] == long_text
    json_content, _, _ = build_selected_export(workbook, "far_records", SelectionManifest(format="json", selection=selection))
    assert json.loads(json_content)["records"][0]["clause_text"] == long_text
    xlsx, _, _ = build_selected_export(workbook, "far_records", SelectionManifest(format="xlsx", selection=selection))
    book = load_workbook(io.BytesIO(xlsx))
    # Excel caps a cell at 32,767 characters: the full text is on Long Text.
    assert "Long Text" in book.sheetnames
    parts = [row[4] for row in book["Long Text"].iter_rows(min_row=5, values_only=True)]
    assert "".join(parts) == long_text
    assert book["Long Text"]["B5"].value == "52.212-4"


@pytest.mark.parametrize(
    "profile_id, dataset_id, values, fields, header",
    [
        (
            "contract_v3",
            "line_items",
            [{"item_number": "0001", "description": "Engineering services", "quantity": "12", "unit": "MO", "unit_price": "$1,000.00", "amount": "$12,000.00"},
             {"item_number": "0002", "description": "Travel", "quantity": "1", "unit": "LOT", "unit_price": "$500.00", "amount": "$500.00"}],
            ["contract.line.quantity", "contract.line.amount"],
            ["Item No.", "Quantity", "Amount"],
        ),
        (
            "invoice",
            "invoice_lines",
            [{"line_number": "1", "description": "Widget", "quantity": "2", "uom": "EA", "unit_price": "5.00", "amount": "10.00"},
             {"line_number": "2", "description": "Gadget", "quantity": "1", "uom": "EA", "unit_price": "7.50", "amount": "7.50"}],
            ["invoice.line.description", "invoice.line.amount"],
            ["Line", "Description", "Line Amount"],
        ),
    ],
)
def test_contract_and_invoice_tables(profile_id, dataset_id, values, fields, header):
    """CLIN / invoice-line tables: selected columns under the document's
    own labels, the identifier (Item No. / Line) carried."""

    workbook = _workbook_for(profile_id, dataset_id, [RawRecord(f"r{i}", v) for i, v in enumerate(values)])
    dataset = next(d for d in workbook.datasets if d.dataset_id == dataset_id)
    known = {c.canonical_field for c in dataset.columns}
    assert set(fields) <= known
    table = select(workbook, dataset_id, SelectionManifest(format="csv", selection=[{"record_id": "r1", "fields": fields}]))
    assert [c.display_label for c in table.columns] == header
    assert len(table.rows) == 1 and table.rows[0][0] == "r1"
    with pytest.raises(SelectionError):
        select(workbook, dataset_id, SelectionManifest(format="csv", selection=[{"record_id": "r9", "fields": fields}]))


def test_transcript_courses():
    from tests.test_academic_transcript import _pdf, _university
    from tests.test_academic_transcript import _workbook as transcript_workbook

    workbook = transcript_workbook(f"transcript-{uuid4()}.pdf", _pdf(_university))
    dataset = next(d for d in workbook["datasets"] if d["cardinality"] == "repeating" and d["records"] and d["role"] == "business")
    record = dataset["records"][0]
    field = next(c["canonical_field"] for c in dataset["columns"] if record["cells"].get(c["canonical_field"], {}).get("value"))
    payload = json.loads(_export(workbook["document_id"], dataset["dataset_id"], "json", [{"record_id": record["record_id"], "fields": [field]}]).content)
    key = next(c["key"] for c in dataset["columns"] if c["canonical_field"] == field)
    assert payload["records"][0][key] == record["cells"][field]["value"]


def test_xml_records():
    from tests.test_xml_document import CLAUSE_XML
    from app.database.session import SessionLocal
    from app.staging.preparation import prepare_staging

    content = CLAUSE_XML.replace("</OKCXMLIMPDFN>", f"<!-- {uuid4()} --></OKCXMLIMPDFN>").encode()
    upload = client.post("/api/documents/upload", files={"file": ("clauses.xml", content, "application/xml")})
    document_id = upload.json()["document_id"]
    client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False})
    database = SessionLocal()
    try:
        prepare_staging(database, database.get(Document, document_id))
    finally:
        database.close()
    workbook = client.get(f"/api/documents/{document_id}/staging-workbook").json()
    dataset = next(d for d in workbook["datasets"] if d["cardinality"] == "repeating" and len(d["records"]) == 2)
    title = next(c for c in dataset["columns"] if c["display_label"] == "Title")
    selection = [{"record_id": r["record_id"], "fields": [title["canonical_field"]]} for r in dataset["records"]]
    rows = _csv(_export(document_id, dataset["dataset_id"], "csv", selection))
    assert rows[0][-1] == "Title"
    assert [r[-1] for r in rows[1:]] == ["FAR 52.201-1 - Acquisition 360: Voluntary Survey", "FAR 52.202-1 - Definitions"]


# --- FAR Final Rule, Part 52, identity, review independence --------------------------------


@pytest.fixture(scope="module")
def final_rule() -> dict:
    from tests.test_far_regulatory_change import RULE_PDF
    from app.database.session import SessionLocal
    from app.staging.preparation import prepare_staging

    if not RULE_PDF.exists():
        pytest.skip("reference Federal Register rule PDF not present")
    content = RULE_PDF.read_bytes() + b"\n%" + uuid4().hex.encode()
    upload = client.post("/api/documents/upload", files={"file": ("2026-04912.pdf", content, "application/pdf")})
    document_id = upload.json()["document_id"]
    client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False})
    database = SessionLocal()
    try:
        assert prepare_staging(database, database.get(Document, document_id)).record.profile_id == "far_regulatory_change"
    finally:
        database.close()
    workbook = client.get(f"/api/documents/{document_id}/staging-workbook").json()
    return {"id": document_id, "datasets": {d["dataset_id"]: d for d in workbook["datasets"]}}


def test_final_rule_data_exports_only_the_chosen_change_fields(final_rule):
    data = final_rule["datasets"]["far_rule_data"]
    changes = [r for r in data["records"] if r["cells"]["rule.data.action"]["value"] == "Replace value"]
    fields = [f"rule.data.{k}" for k in ("far_reference", "action", "previous_value", "new_value", "effective_date", "fac_number", "far_case")]
    selection = [{"record_id": r["record_id"], "fields": fields} for r in changes]
    rows = _csv(_export(final_rule["id"], "far_rule_data", "csv", selection))
    assert rows[0] == ["FAR Reference", "Action / Change", "Previous Value", "New Value", "Effective Date", "FAC Number", "FAR Case"]
    assert ["22.1503(b)(2)", "Replace value", "$102,280", "$105,767", "March 13, 2026", "2026-01", "2025-007"] in rows[1:]
    assert len(rows) - 1 == len(changes)


def test_final_rule_summary_fields(final_rule):
    summary = final_rule["datasets"]["rule_summary"]["records"][0]
    selection = [{"record_id": summary["record_id"], "fields": ["rule.heading", "rule.reference", "rule.effective_date"]}]
    payload = json.loads(_export(final_rule["id"], "rule_summary", "json", selection).content)
    assert payload["records"] == [{
        "heading": "Federal Acquisition Regulation — Trade Agreements Thresholds",
        "reference": "FAC 2026-01 | FAR Case 2025-007",
        "effective_date": "March 13, 2026",
    }]
    assert [f["canonical_field"] for f in payload["fields"]] == ["rule.heading", "rule.reference", "rule.effective_date"]


def test_part_52_clause_text_and_json_identity():
    from tests.test_far_part_52 import _prepare, _upload, far_html

    document_id = _upload(far_html(), "part_52_selection.html")
    assert _prepare(document_id)["profile_id"] == "far_part_52"
    records = next(
        d for d in client.get(f"/api/documents/{document_id}/staging-workbook").json()["datasets"] if d["dataset_id"] == "far_records"
    )["records"]
    clause = next(r for r in records if r["cells"][F + "far_number"]["value"] == "52.212-5")
    payload = json.loads(_export(document_id, "far_records", "json", [{"record_id": clause["record_id"], "fields": [F + "clause_text"]}]).content)
    assert payload["identifier"] == "far_number"
    assert payload["records"][0]["far_number"] == "52.212-5"
    assert payload["records"][0]["clause_text"] == clause["cells"][F + "clause_text"]["value"]
    assert "Contract Terms and Conditions" in payload["records"][0]["clause_text"]


def test_selection_is_not_verification(far):
    """Exporting a selection never changes Verified / Needs Review / Missing."""

    def statuses():
        workbook = client.get(f"/api/documents/{far['id']}/staging-workbook").json()
        return [
            (r["record_id"], r["record_status"], tuple(c["review_status"] for c in r["cells"].values()))
            for d in workbook["datasets"] for r in d["records"]
        ]

    before = statuses()
    for fmt in ("xlsx", "csv", "json"):
        assert _export(far["id"], "far_records", fmt, [{"record_id": _rid(far, "46.101"), "fields": [F + "title"]}]).status_code == 200
    assert statuses() == before


# --- label / value lists (Supplier, Charges & Totals) ---------------------------------------


def _invoice_parts_workbook():
    profile = registry.get_profile("invoice")
    document = Document(id="doc-fields", original_filename="Invoice 77.pdf", stored_filename="doc-fields.pdf")
    return assemble_workbook(
        profile=profile,
        document=document,
        adapter_result=AdapterResult(
            records={
                "supplier": [RawRecord("supplier", {"name": "Summit Industrial", "email": "ar@summit.test"})],
                "taxes_charges": [RawRecord("c1", {"label": "Fuel Surcharge", "amount": "74.00"})],
            }
        ),
        metadata=ProcessingMetadata(),
    )


def test_selected_fields_read_labels_and_values_from_the_workbook():
    from app.staging.selection_export import FieldSelectionManifest, build_selected_fields_export, select_fields

    workbook = _invoice_parts_workbook()
    manifest = FieldSelectionManifest(
        format="csv",
        fields=[
            {"dataset_id": "taxes_charges", "record_id": "c1", "field": "invoice.charge.amount"},
            {"dataset_id": "supplier", "record_id": "supplier", "field": "invoice.supplier.email"},
            {"dataset_id": "supplier", "record_id": "supplier", "field": "invoice.supplier.email"},  # once only
        ],
    )
    table = select_fields(workbook, manifest)
    sections = {d.dataset_id: d.display_name for d in workbook.datasets}
    # The reader's order, each value named as it reads in the list.
    assert [(v["section"], v["field"]) for _, v in table.rows] == [
        (sections["taxes_charges"], "Fuel Surcharge"),
        (sections["supplier"], "Email"),
    ]
    assert table.rows[1][1]["value"] == "ar@summit.test"
    content, _, filename = build_selected_fields_export(workbook, manifest)
    assert list(csv.reader(io.StringIO(content)))[0] == ["Section", "Field", "Value"]
    assert filename == "Invoice 77_selected_fields.csv"
    for fmt in ("json", "xlsx"):
        build_selected_fields_export(workbook, FieldSelectionManifest(format=fmt, fields=manifest.fields))


@pytest.mark.parametrize(
    "item",
    [
        {"dataset_id": "nope", "record_id": "supplier", "field": "invoice.supplier.email"},
        {"dataset_id": "supplier", "record_id": "other", "field": "invoice.supplier.email"},
        {"dataset_id": "supplier", "record_id": "supplier", "field": "invoice.charge.amount"},
    ],
)
def test_selected_fields_reject_what_the_document_does_not_hold(item):
    from app.staging.selection_export import FieldSelectionManifest, select_fields

    with pytest.raises(SelectionError):
        select_fields(_invoice_parts_workbook(), FieldSelectionManifest(format="csv", fields=[item]))


def test_selected_fields_endpoint(far):
    record = far["records"][0]
    field = F + "far_number"
    response = client.post(
        f"/api/documents/{far['id']}/staging-workbook/export-selected-fields",
        json={"format": "csv", "fields": [{"dataset_id": "far_records", "record_id": record["record_id"], "field": field}]},
    )
    rows = _csv(response)
    assert rows[0] == ["Section", "Field", "Value"] and rows[1][2] == record["cells"][field]["value"]
    bad = client.post(
        f"/api/documents/{far['id']}/staging-workbook/export-selected-fields",
        json={"format": "csv", "fields": [{"dataset_id": "far_records", "record_id": "missing", "field": field}]},
    )
    assert bad.status_code == 422
