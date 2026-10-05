"""FAR_REGULATION — one rule for any FAR Part (1–53): regulatory function as
Record Type, prescriptions (51.107 → 52.251-1), typed cross references,
Part / Subpart titles, Reserved records, Part 52 clause semantics only for
Part 52, and Part 53 form fields."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from app.far.regulation import classify_record, form_info, part_title, prescriptions, references_text, typed_references
from app.main import app
from tests.test_far_part_52 import _article, _prepare, _upload

client = TestClient(app)
PART_51 = Path(__file__).resolve().parents[2] / "reference" / "regression" / "far_part_51.pdf"
R = "far.record."


# --- the rule's parts -----------------------------------------------------------------------


def test_prescriptions_name_what_a_section_prescribes():
    p = prescriptions("The contracting officer shall insert the clause at 52.251-1, Government Supply Sources, in solicitations and contracts when …")
    assert p.references == ["52.251-1"] and p.kind == "Contract Clause Prescription"
    p = prescriptions("The contracting officer shall insert the provision at 52.216-1, Type of Contract, in solicitations …")
    assert p.kind == "Solicitation Provision Prescription"
    p = prescriptions(
        "(a) The contracting officer shall insert the provision at 52.237-1, Site Visit, in solicitations for services.\n"
        "(b) The contracting officer shall insert the clause at 52.237-2, Protection of Government Buildings, …\n"
        "(c) The contracting officer shall insert the clause at 52.237-3, Continuity of Services, …"
    )
    assert p.references == ["52.237-1", "52.237-2", "52.237-3"] and p.kind == "Provision & Clause Prescription"
    assert len(p.usage.split("\n")) == 3
    # A mention is not a prescription.
    assert prescriptions("See 52.216-1 for the type of contract.") is None


def test_record_type_is_the_regulatory_function():
    assert classify_record("Definitions", "As used in this part— Personal services contract means …", None) == "Definition"
    assert classify_record("Scope of part", "This part prescribes …", None) == "Scope"
    assert classify_record("Policy", "Agencies shall …", None) == "Policy"
    assert classify_record("Contracting officer responsibility", "The contracting officer is responsible …", None) == "Responsibility"
    assert classify_record("Limitations", "…", None) == "Limitation"
    assert classify_record("[Reserved]", "", None, reserved=True) == "Reserved"
    # The text decides before the heading: a definitions list under any heading.
    assert classify_record("General", "As used in this subpart— Award fee means …", None) == "Definition"
    p = prescriptions("The contracting officer shall insert the clause at 52.251-1 …")
    assert classify_record("Contract clause", "…", p) == "Contract Clause Prescription"


def test_typed_cross_references():
    text = (
        "Insert 52.216-1 (see subpart 37.6 and part 35, and 15.404-1(b)). Under 41 U.S.C. 3902 and 5 CFR Part 300, "
        "Pub. L. 115-232, E.O. 12866 and Standard Form 1449 …"
    )
    refs = typed_references(text, own_number="16.105")
    assert [(r.kind, r.text) for r in refs] == [
        ("FAR", "52.216-1"), ("FAR", "subpart 37.6"), ("FAR", "part 35"), ("FAR", "15.404-1(b)"),
        ("USC", "41 U.S.C. 3902"), ("CFR", "5 CFR Part 300"), ("Public Law", "Pub. L. 115-232"),
        ("Executive Order", "E.O. 12866"), ("Form", "Standard Form 1449"),
    ]
    # One column, FAR first.
    assert references_text(refs) == (
        "52.216-1; subpart 37.6; part 35; 15.404-1(b); 5 CFR Part 300; 41 U.S.C. 3902; Pub. L. 115-232; "
        "E.O. 12866; Standard Form 1449"
    )
    # A record never cites itself.
    assert typed_references("This section 16.105 applies.", own_number="16.105") == []


def test_part_53_form_fields():
    info = form_info(
        "Standard Form 1449, Solicitation/Contract/Order for Commercial Products and Commercial Services",
        "SF 1449 is prescribed in 12.204(a). It supersedes the 2021 edition.",
    )
    assert (info.number, info.form_type, info.name) == (
        "SF 1449", "Standard Form", "Solicitation/Contract/Order for Commercial Products and Commercial Services",
    )
    assert info.prescribing_reference == "12.204(a)" and info.supersession == "It supersedes the 2021 edition."
    assert form_info("Policy", "…") is None
    assert part_title("Part 15 - Contracting by Negotiation") == "Contracting by Negotiation"


# --- a real Part (Parts 1–51 schema) -------------------------------------------------------


def _records(document_id: str) -> tuple[dict, dict[str, dict]]:
    workbook = client.get(f"/api/documents/{document_id}/staging-workbook").json()
    dataset = next(d for d in workbook["datasets"] if d["dataset_id"] == "far_records")
    rows = {
        r["cells"][R + "far_number"]["value"]: {k.removeprefix(R): c["value"] for k, c in r["cells"].items()}
        for r in dataset["records"]
    }
    return dataset, rows


@pytest.fixture(scope="module")
def part_51() -> str:
    if not PART_51.exists():
        pytest.skip("reference Part 51 PDF not present")
    content = PART_51.read_bytes() + b"\n%" + uuid4().hex.encode()
    upload = client.post("/api/documents/upload", files={"file": ("Part-51.pdf", content, "application/pdf")})
    document_id = upload.json()["document_id"]
    assert client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False}).status_code == 200
    assert _prepare(document_id)["profile_id"] == "far_part_52"
    return document_id


def test_part_51_prescription_and_hierarchy(part_51):
    dataset, rows = _records(part_51)
    assert dataset["display_name"] == "FAR Data"
    labels = {c["canonical_field"]: c["display_label"] for c in dataset["columns"]}
    assert labels[R + "section_text"] == "Description / Regulatory Text"
    clause = rows["51.107"]
    assert clause["title"] == "Contract clause"
    assert clause["record_type"] == "Contract Clause Prescription"
    assert (clause["far_part"], clause["far_part_title"]) == ("Part 51", "Use of Government Sources by Contractors")
    assert (clause["far_subpart"], clause["far_subpart_title"]) == ("Subpart 51.1", "Contractor Use of Government Supply Sources")
    assert clause["prescription_reference"] == "52.251-1"
    assert "52.251-1" in clause["prescription"] and "52.251-1" in clause["cross_references"]
    assert rows["51.000"]["record_type"] == "Scope" and rows["51.101"]["record_type"] == "Policy"
    # Part 52 fields are never forced onto Part 51.
    for record in rows.values():
        assert record["clause_text"] is None and record["provision_text"] is None and record["revision_date"] is None


def test_part_51_exports_use_the_core_schema(part_51):
    base = f"/api/documents/{part_51}/staging-workbook/exports"
    rows = list(csv.reader(io.StringIO(client.get(f"{base}/far_clauses_provisions.csv").content.decode("utf-8-sig"))))
    assert rows[0] == [
        "FAR Number", "Title", "Record Type", "FAR Part", "FAR Part Title", "FAR Subpart", "FAR Subpart Title",
        "FAR Section", "Record Status", "Description / Regulatory Text", "Prescription / Usage",
        "Prescription Reference", "Cross References", "Source Reference",
    ]
    payload = json.loads(client.get(f"{base}/far_canonical.json").content)
    cited = {r["far_number"]: r["references"] for r in payload["references"]}
    assert {"type": "FAR", "reference": "52.251-1"} in cited["51.107"]
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(client.get(f"{base}/far_workbook.xlsx").content))
    assert book.sheetnames[0] == "FAR DATA"
    assert book["FAR DATA"]["A1"].value == "FAR Part 51 — FAR Data"


# --- Part 53 (Forms) -------------------------------------------------------------------------


def part_53_html() -> bytes:
    def section(number: str, title: str, body: str) -> str:
        return _article("FAR_" + number.replace(".", "_").replace("-", "_"), 3, number, f"{title}.", f'<p class="p">{body}</p>')

    filler = "".join(section(f"53.10{i}", f"Requirements {i}", "Agencies shall use the forms as prescribed.") for i in range(1, 8))
    sections = (
        section("53.000", "Scope of part", "This part prescribes standard forms and references optional forms.")
        + filler
        + section("53.212", "Acquisition of commercial products and commercial services (SF 1449)",
                  "SF 1449, prescribed in 12.204(a), is used for commercial acquisitions.")
        + section("53.301-1449", "Standard Form 1449, Solicitation/Contract/Order for Commercial Products and Commercial Services",
                  "This form is prescribed in 12.204(a). Previous editions are obsolete.")
        + section("53.302-347", "Optional Form 347, Order for Supplies or Services", "Prescribed in 13.307(b).")
    )
    subparts = (
        _article("FAR_Subpart_53_1", 2, "Subpart 53.1", "General", None, sections[: len(sections) // 2])
        + _article("FAR_Subpart_53_2", 2, "Subpart 53.2", "Prescription of Forms", None, sections[len(sections) // 2 :])
        + _article("FAR_Subpart_53_3", 2, "Subpart 53.3", "Illustrations of Forms", None, "")
    )
    part = _article("FAR_Part_53", 1, "Part 53", "Forms", None, subparts)
    html = f"<html><body><main>{part}<p>Parent topic: Federal Acquisition Regulation</p></main></body></html><!-- {uuid4()} -->"
    return html.encode()


def test_part_53_records_carry_form_fields():
    document_id = _upload(part_53_html(), "part_53.html")
    assert _prepare(document_id)["profile_id"] == "far_part_52"
    _, rows = _records(document_id)
    sf = rows["53.301-1449"]
    assert (sf["form_number"], sf["form_type"]) == ("SF 1449", "Standard Form")
    assert sf["form_name"] == "Solicitation/Contract/Order for Commercial Products and Commercial Services"
    assert sf["prescribing_reference"] == "12.204(a)" and sf["supersession"] == "Previous editions are obsolete."
    assert rows["53.302-347"]["form_number"] == "OF 347"
    assert rows["53.212"]["form_number"] == "SF 1449"
    assert rows["53.000"]["form_number"] is None
    header = next(csv.reader(io.StringIO(
        client.get(f"/api/documents/{document_id}/staging-workbook/exports/far_clauses_provisions.csv").content.decode("utf-8-sig")
    )))
    assert "Form Number" in header and "Clause Text" not in header
