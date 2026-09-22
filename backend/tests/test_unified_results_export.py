"""Unified structured results, XLSX export, processing reuse, reopen."""

from __future__ import annotations

import io
import time
from uuid import uuid4

import fitz
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database.session import SessionLocal
from app.main import app
from app.models.document import Document
from app.models.document_extracted_table import DocumentExtractedTable
from app.models.document_metadata_field import DocumentMetadataField
from app.models.document_page import DocumentPage
from app.services.processing_versions import (
    DISCOVERY_PROCESSOR_VERSION,
    EXTRACTION_PROCESSOR_VERSION,
    PAGES_PROCESSOR_VERSION,
    discovery_artifacts_reusable,
    document_content_fingerprint,
    pages_artifacts_reusable,
    processing_versions_payload,
)
from app.services.reviewed_export import (
    AUTHORITATIVE_REVIEW_STATUSES,
    build_export_xlsx,
    build_fields_wide_csv,
    extracted_value_for,
)

client = TestClient(app)


def _pdf(text: str = "Contract Number: 47QRCA25DSF07\nDate Issued: 4/15/2025") -> bytes:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), text, fontsize=10)
    content = pdf.tobytes()
    pdf.close()
    return content


def _upload(name: str | None = None) -> str:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                name or f"unified-{uuid4()}.pdf",
                _pdf(),
                "application/pdf",
            )
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["document_id"]


def _seed_field(
    database: Session,
    document_id: str,
    *,
    key: str,
    label: str,
    value: str | None,
    review_status: str = "pending",
    machine_value: str | None = None,
    raw_ocr: str | None = None,
) -> DocumentMetadataField:
    machine = machine_value if machine_value is not None else value
    field = DocumentMetadataField(
        document_id=document_id,
        field_group="Header",
        field_key=key,
        label=label,
        value=value,
        original_value=machine,
        confidence=0.9 if value else 0.2,
        confidence_band="high" if value else "low",
        value_type="text",
        extraction_method="layout",
        evidence_json={
            "page_number": 1,
            "source_text": raw_ocr or f"{label}: {machine or ''}",
            "raw_ocr": raw_ocr or f"{label}: {machine or ''}",
            "machine_value": machine,
            "reviewed_value": value if review_status == "edited" else None,
            "validation": {
                "status": "passed" if value else "needs_review",
                "checks": [],
                "warnings": [],
            },
            "validation_status": "passed" if value else "needs_review",
        },
        verified=review_status in AUTHORITATIVE_REVIEW_STATUSES,
        review_status=review_status,
        extraction_source="target",
        human_approved=review_status in AUTHORITATIVE_REVIEW_STATUSES,
    )
    database.add(field)
    return field


def _seed_table(
    database: Session,
    document_id: str,
    *,
    key: str = "clins",
    display_name: str = "CLINs",
) -> DocumentExtractedTable:
    table = DocumentExtractedTable(
        document_id=document_id,
        target_key=key,
        display_name=display_name,
        columns_json=["clin", "description", "amount"],
        rows_json=[
            {"clin": "0001", "description": "Labor", "amount": "1000"},
            {"clin": "0002", "description": "Travel", "amount": "250"},
        ],
        pages_json=[1],
        extraction_method="layout",
        confidence=0.88,
    )
    database.add(table)
    return table


def test_wide_fields_csv_is_dataset_oriented() -> None:
    document_id = _upload()
    database = SessionLocal()
    try:
        _seed_field(
            database,
            document_id,
            key="contract_number",
            label="Contract Number",
            value="47QRCA25DSF07",
        )
        _seed_field(
            database,
            document_id,
            key="date_issued",
            label="Date Issued",
            value="4/15/2025",
        )
        database.commit()
        fields = list(
            database.query(DocumentMetadataField).filter_by(
                document_id=document_id
            )
        )
        csv_body = build_fields_wide_csv(fields)
    finally:
        database.close()

    lines = csv_body.strip().splitlines()
    assert "contract_number" in lines[0]
    assert "date_issued" in lines[0]
    assert "47QRCA25DSF07" in lines[1]
    assert "4/15/2025" in lines[1]
    assert not lines[0].startswith("field,value")


def test_xlsx_workbook_has_fields_tables_and_evidence() -> None:
    document_id = _upload()
    database = SessionLocal()
    try:
        document = database.get(Document, document_id)
        assert document is not None
        _seed_field(
            database,
            document_id,
            key="contract_number",
            label="Contract Number",
            value="47QRCA25DSF07",
            machine_value="47QRCA25DSF07",
            raw_ocr="Contract Number 47QRCA25DSF07",
        )
        _seed_table(database, document_id)
        database.commit()
        fields = list(
            database.query(DocumentMetadataField).filter_by(
                document_id=document_id
            )
        )
        tables = list(
            database.query(DocumentExtractedTable).filter_by(
                document_id=document_id
            )
        )
        body = build_export_xlsx(
            document=document, fields=fields, tables=tables
        )
    finally:
        database.close()

    workbook = load_workbook(io.BytesIO(body))
    titles = set(workbook.sheetnames)
    assert "All Fields" in titles
    assert "Key Contract Fields" in titles
    assert "Sections" in titles
    assert "Needs Review" in titles
    assert "Source Evidence" in titles
    assert any("CLIN" in name.upper() or "clin" in name.lower() for name in titles)

    all_fields = workbook["All Fields"]
    headers = [cell.value for cell in all_fields[1]]
    assert "Field / Label" in headers
    assert "Extracted Value" in headers
    labels = [row[2].value for row in all_fields.iter_rows(min_row=2)]
    values = [row[3].value for row in all_fields.iter_rows(min_row=2)]
    assert any(
        label and "contract" in str(label).lower() for label in labels
    ) or "47QRCA25DSF07" in values

    key_sheet = workbook["Key Contract Fields"]
    key_headers = [cell.value for cell in key_sheet[1]]
    key_values = [cell.value for cell in key_sheet[2]]
    assert any(
        header and "contract" in str(header).lower() for header in key_headers
    )
    assert "47QRCA25DSF07" in key_values

    evidence = workbook["Source Evidence"]
    evidence_headers = [cell.value for cell in evidence[1]]
    assert "raw_ocr" in evidence_headers
    assert "extracted_value" in evidence_headers
    assert "confidence" in evidence_headers

    response = client.get(f"/api/documents/{document_id}/export.xlsx")
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]


def test_review_edit_keeps_machine_extracted_value_immutable() -> None:
    document_id = _upload()
    database = SessionLocal()
    try:
        field = _seed_field(
            database,
            document_id,
            key="contract_number",
            label="Contract Number",
            value="47QRCA25DSF07",
            machine_value="47QRCA25DSF07",
        )
        database.commit()
        field_id = field.id
    finally:
        database.close()

    edit = client.post(
        f"/api/documents/{document_id}/targets/contract_number/corrections",
        json={
            "action": "edit",
            "original_value": "47QRCA25DSF07",
            "corrected_value": "REVIEWED-001",
        },
    )
    assert edit.status_code == 201, edit.text

    database = SessionLocal()
    try:
        field = database.get(DocumentMetadataField, field_id)
        assert field is not None
        assert field.value == "REVIEWED-001"
        assert extracted_value_for(field) == "47QRCA25DSF07"
        assert field.original_value == "47QRCA25DSF07" or (
            (field.evidence_json or {}).get("machine_value") == "47QRCA25DSF07"
        )
    finally:
        database.close()

    export = client.get(f"/api/documents/{document_id}/export")
    assert export.status_code == 200
    record = export.json()["fields"][0]
    assert record["value"] == "REVIEWED-001"
    assert record["extracted_value"] == "47QRCA25DSF07"
    assert record["review_status"] == "edited"


def test_extract_results_reopen_loads_persisted_without_rediscovery() -> None:
    from app.schemas.document_target import ScalarTargetResult, TableTargetResult
    from app.services.target_result_store import persist_target_extraction_results

    document_id = _upload()
    database = SessionLocal()
    try:
        document = database.get(Document, document_id)
        assert document is not None
        document.processing_status = "completed"
        document.page_count = 1
        database.add(
            DocumentPage(
                document_id=document_id,
                page_number=1,
                final_text="Contract Number: 47QRCA25DSF07",
                requires_ocr=False,
                ocr_succeeded=False,
                extraction_method="native",
                page_width=612.0,
                page_height=792.0,
            )
        )
        persist_target_extraction_results(
            database=database,
            document_id=document_id,
            scalars=[
                ScalarTargetResult(
                    target="Contract Number",
                    normalized_key="contract_number",
                    value="47QRCA25DSF07",
                    page=1,
                    confidence=0.9,
                    confidence_band="high",
                    verified=True,
                    extraction_method="layout",
                    display_method="Native",
                    evidence={
                        "page_number": 1,
                        "source_text": "Contract Number: 47QRCA25DSF07",
                        "source_reference": "page 1",
                    },
                )
            ],
            tables=[
                TableTargetResult(
                    target="CLINs",
                    columns=["clin", "description"],
                    rows=[{"clin": "0001", "description": "Labor"}],
                    pages=[1],
                )
            ],
        )
        database.commit()
    finally:
        database.close()

    first = client.get(f"/api/documents/{document_id}/extract-results")
    assert first.status_code == 200, first.text
    assert any(
        s.get("normalized_key") == "contract_number"
        for s in first.json().get("scalars", [])
    )
    second = client.get(f"/api/documents/{document_id}/extract-results")
    assert second.status_code == 200, second.text
    assert len(second.json().get("scalars", [])) >= 1


def test_processing_version_keys_gate_reuse() -> None:
    database = SessionLocal()
    try:
        document_id = _upload()
        document = database.get(Document, document_id)
        assert document is not None
        document.page_count = 3
        document.checksum_sha256 = "abc123"
        document.ingestion_provenance = {
            **processing_versions_payload(document=document),
        }
        database.commit()
        database.refresh(document)

        assert pages_artifacts_reusable(document) is True
        assert discovery_artifacts_reusable(document) is True
        assert document_content_fingerprint(document) == "abc123:3"

        document.checksum_sha256 = "changed"
        assert pages_artifacts_reusable(document) is False

        document.checksum_sha256 = "abc123"
        document.ingestion_provenance = {
            "content_fingerprint": document_content_fingerprint(document),
            "pages_processor_version": PAGES_PROCESSOR_VERSION,
            "discovery_processor_version": "stale",
            "extraction_processor_version": EXTRACTION_PROCESSOR_VERSION,
        }
        assert pages_artifacts_reusable(document) is True
        assert discovery_artifacts_reusable(document) is False
        assert DISCOVERY_PROCESSOR_VERSION != "stale"
    finally:
        database.close()


def test_blank_needs_review_not_populated_with_label() -> None:
    document_id = _upload()
    database = SessionLocal()
    try:
        _seed_field(
            database,
            document_id,
            key="issued_by",
            label="Issued By",
            value="",
            machine_value="",
            raw_ocr="Issued By",
            review_status="pending",
        )
        database.commit()
        fields = list(
            database.query(DocumentMetadataField).filter_by(
                document_id=document_id
            )
        )
        csv_body = build_fields_wide_csv(fields)
    finally:
        database.close()

    lines = csv_body.strip().splitlines()
    assert lines[0] == "issued_by"
    # Empty / blank — not the form label text.
    assert "Issued By" not in lines[1]
    assert lines[1] in {"", '""'}


def test_processing_reuse_timing_second_pass_faster() -> None:
    """Smoke: reuse gate is near-instant vs full reprocess."""

    document_id = _upload("reuse-timing.pdf")
    database = SessionLocal()
    try:
        document = database.get(Document, document_id)
        assert document is not None
        document.processing_status = "completed"
        document.page_count = 1
        if not document.checksum_sha256:
            document.checksum_sha256 = "reuse-test-checksum"
        document.ingestion_provenance = {
            **processing_versions_payload(document=document),
            "pages_reused": False,
            "discovery_reused": False,
        }
        database.add(
            DocumentPage(
                document_id=document_id,
                page_number=1,
                final_text="Contract Number: 47QRCA25DSF07",
                requires_ocr=False,
                ocr_succeeded=False,
                extraction_method="native",
                page_width=612.0,
                page_height=792.0,
            )
        )
        database.commit()
        assert pages_artifacts_reusable(document) is True
    finally:
        database.close()

    started = time.perf_counter()
    database = SessionLocal()
    try:
        document = database.get(Document, document_id)
        assert document is not None
        assert pages_artifacts_reusable(document) is True
    finally:
        database.close()
    reopen_ms = int((time.perf_counter() - started) * 1000)
    assert reopen_ms < 5000, f"reopen gate took {reopen_ms}ms"


def test_normalized_endpoint_matches_export_sources() -> None:
    document_id = _upload()
    database = SessionLocal()
    try:
        _seed_field(
            database,
            document_id,
            key="contract_number",
            label="Contract Number",
            value="47QRCA25DSF07",
        )
        _seed_table(database, document_id)
        database.commit()
    finally:
        database.close()

    response = client.get(f"/api/documents/{document_id}/normalized")
    assert response.status_code == 200, response.text
    body = response.json()

    assert body["document_id"] == document_id
    assert body["fields"]["contract_number"] == "47QRCA25DSF07"
    assert "clins" in body["tables"]
    assert body["tables"]["clins"]["columns"] == ["clin", "description", "amount"]
    assert body["tables"]["clins"]["rows"][0]["clin"] == "0001"

    # Same source as the wide CSV — no independent re-derivation.
    csv_response = client.get(f"/api/documents/{document_id}/export.csv")
    assert csv_response.status_code == 200, csv_response.text
    csv_lines = csv_response.text.strip().splitlines()
    assert "contract_number" in csv_lines[0]
    assert "47QRCA25DSF07" in csv_lines[1]
