"""Universal Document Ingestion RC — certification matrix (minimum)."""

from __future__ import annotations

from io import BytesIO
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile, ZipInfo

import fitz
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation

from app.main import app
from app.services.libreoffice_convert import LEGACY_OFFICE_UNAVAILABLE
from app.services.security_validation import detect_kind_from_bytes


client = TestClient(app)


def _pdf_bytes(pages: int = 1, text: str | None = None) -> bytes:
    pdf = fitz.open()
    for index in range(pages):
        page = pdf.new_page()
        page.insert_text(
            (72, 72),
            text or f"Native page {index + 1} {uuid4()}",
        )
    data = pdf.tobytes()
    pdf.close()
    return data


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Matrix DOCX")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = str(uuid4())
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "One"
    first["A1"] = "hello"
    second = workbook.create_sheet("Two")
    second["A1"] = f"sheet-two-{uuid4()}"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = f"Slide {uuid4()}"
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _png_bytes() -> bytes:
    image = Image.new("RGB", (240, 80), color="white")
    ImageDraw.Draw(image).text((8, 30), str(uuid4()), fill="black")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_ready_exposes_legacy_office_conversion() -> None:
    response = client.get("/ready")
    assert response.status_code == 200
    body = response.json()
    assert "legacy_office_conversion" in body
    assert "available" in body["legacy_office_conversion"]
    assert "legacy_office_conversion" in body["checks"]


def test_signature_wins_over_pdf_extension() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "not-really.pdf",
                _png_bytes(),
                "application/pdf",
            )
        },
    )
    assert response.status_code == 400
    assert "signature" in response.json()["detail"].lower() or "Detected" in response.json()["detail"]


def test_txt_rejects_pdf_bytes() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": ("notes.txt", _pdf_bytes(), "text/plain"),
        },
    )
    assert response.status_code == 400


def test_detect_kind_helpers() -> None:
    assert detect_kind_from_bytes(b"%PDF-1.7") == "pdf"
    assert detect_kind_from_bytes(b"\xff\xd8\xff\xe0") == "jpeg"


def test_one_page_pdf_provenance_and_pages() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": ("one.pdf", _pdf_bytes(1), "application/pdf"),
        },
    )
    assert response.status_code == 201
    body = response.json()
    assert body["size_tier"] == "small"
    assert body["prefer_background"] is False
    assert body["ingestion_provenance"]["source_format"] == "pdf"

    document_id = body["document_id"]
    extraction = client.post(
        f"/api/documents/{document_id}/extract-pages",
        json={
            "run_ocr": True,
            "page_start": None,
            "page_end": None,
            "force_reprocess": False,
        },
    )
    assert extraction.status_code == 200
    pages = client.get(f"/api/documents/{document_id}/pages")
    assert pages.status_code == 200
    assert len(pages.json()) == 1
    assert "Native page" in pages.json()[0]["final_text"]


def test_multi_sheet_xlsx_page_count() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "book.xlsx",
                _xlsx_bytes(),
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"
                ),
            )
        },
    )
    assert response.status_code == 201
    assert response.json()["page_count"] == 2
    pages = client.get(
        f"/api/documents/{response.json()['document_id']}/pages"
    )
    labels = {page.get("page_label") for page in pages.json()}
    assert "One" in labels
    assert "Two" in labels


def test_docx_tables_preserved() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "table.docx",
                _docx_bytes(),
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".wordprocessingml.document"
                ),
            )
        },
    )
    assert response.status_code == 201
    pages = client.get(
        f"/api/documents/{response.json()['document_id']}/pages"
    )
    page = pages.json()[0]
    assert "Matrix DOCX" in page["final_text"]
    assert "A | B" in page["final_text"] or "1 |" in page["final_text"]


def test_pptx_and_png_upload() -> None:
    pptx = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "deck.pptx",
                _pptx_bytes(),
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".presentationml.presentation"
                ),
            )
        },
    )
    assert pptx.status_code == 201
    assert pptx.json()["page_count"] >= 1
    pages = client.get(f"/api/documents/{pptx.json()['document_id']}/pages")
    assert "Slide" in pages.json()[0]["final_text"]

    png = client.post(
        "/api/documents/upload",
        files={"file": ("scan.png", _png_bytes(), "image/png")},
    )
    assert png.status_code == 201
    assert png.json()["ingestion_provenance"]["source_format"] == "png"


def test_legacy_office_missing_soffice_fails_gracefully() -> None:
    # Minimal OLE header — enough to pass signature, fail conversion.
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    with patch(
        "app.services.upload_processing.libreoffice_available",
        return_value=False,
    ):
        response = client.post(
            "/api/documents/upload",
            files={
                "file": ("legacy.doc", ole, "application/msword"),
            },
        )
    assert response.status_code == 400
    assert response.json()["detail"] == LEGACY_OFFICE_UNAVAILABLE


def test_medium_upload_enqueues_background_job() -> None:
    # Force medium tier without creating a 20MB fixture on disk.
    with patch(
        "app.services.upload_processing.prefer_background_processing",
        return_value=True,
    ), patch(
        "app.services.upload_processing.classify_upload_size",
        return_value="medium",
    ):
        response = client.post(
            "/api/documents/upload",
            files={
                "file": ("med.pdf", _pdf_bytes(2), "application/pdf"),
            },
        )
    assert response.status_code == 201
    body = response.json()
    assert body["prefer_background"] is True
    assert body["size_tier"] == "medium"
    assert body["processing_job_id"] is not None


def test_corrupt_ooxml_rejected() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr(ZipInfo("readme.txt"), "no workbook")
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "bad.xlsx",
                buffer.getvalue(),
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"
                ),
            )
        },
    )
    assert response.status_code == 400
