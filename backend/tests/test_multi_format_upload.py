"""Multi-format upload + native ingest regression tests."""

from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile, ZipInfo

from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation

from app.main import app
from app.services.ocr_detection import suspicious_glyph_ratio
from app.services.upload_size_tiers import (
    classify_upload_size,
    prefer_background_processing,
)


client = TestClient(app)


def create_test_docx() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Master Services Agreement")
    document.add_paragraph(
        f"This agreement is between Acme and Globex. {uuid4()}"
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def create_test_png() -> bytes:
    image = Image.new("RGB", (300, 80), color="white")
    ImageDraw.Draw(image).text((10, 30), str(uuid4()), fill="black")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def create_test_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet["A1"] = "Item"
    sheet["B1"] = "Amount"
    sheet["A2"] = f"Widget-{uuid4()}"
    sheet["B2"] = 42
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def create_test_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )
    slide.shapes.title.text = f"Briefing {uuid4()}"
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_upload_docx_extracts_text_immediately() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "msa.docx",
                create_test_docx(),
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".wordprocessingml.document"
                ),
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "ready"
    assert body["page_count"] == 1

    pages = client.get(f"/api/documents/{body['document_id']}/pages")
    assert pages.status_code == 200
    assert "Master Services Agreement" in pages.json()[0]["final_text"]


def test_reject_docx_that_is_not_really_ooxml() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "fake.docx",
                b"not actually a docx",
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".wordprocessingml.document"
                ),
            )
        },
    )
    assert response.status_code == 400


def test_upload_png_and_extract_pages() -> None:
    response = client.post(
        "/api/documents/upload",
        files={"file": ("scan.png", create_test_png(), "image/png")},
    )
    assert response.status_code == 201
    body = response.json()
    assert body["content_type"] == "image/png"

    extraction = client.post(
        f"/api/documents/{body['document_id']}/extract-pages",
        json={
            "run_ocr": True,
            "page_start": None,
            "page_end": None,
            "force_reprocess": False,
        },
    )
    assert extraction.status_code == 200
    assert extraction.json()["pages_processed"] == 1


def test_upload_txt_native_ingest() -> None:
    payload = f"Contract notice {uuid4()}\nClause 1 applies.\n".encode()
    response = client.post(
        "/api/documents/upload",
        files={"file": ("notes.txt", payload, "text/plain")},
    )
    assert response.status_code == 201
    document_id = response.json()["document_id"]
    pages = client.get(f"/api/documents/{document_id}/pages")
    assert pages.status_code == 200
    assert "Contract notice" in pages.json()[0]["final_text"]


def test_upload_csv_native_ingest() -> None:
    payload = f"name,value\nalpha,{uuid4()}\n".encode()
    response = client.post(
        "/api/documents/upload",
        files={"file": ("rates.csv", payload, "text/csv")},
    )
    assert response.status_code == 201
    pages = client.get(
        f"/api/documents/{response.json()['document_id']}/pages"
    )
    assert "alpha" in pages.json()[0]["final_text"]


def test_upload_xlsx_sheets_as_pages() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "book.xlsx",
                create_test_xlsx(),
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"
                ),
            )
        },
    )
    assert response.status_code == 201
    body = response.json()
    assert body["page_count"] >= 1
    pages = client.get(f"/api/documents/{body['document_id']}/pages")
    assert pages.status_code == 200
    assert "Widget-" in pages.json()[0]["final_text"]


def test_upload_pptx_slides_as_pages() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "deck.pptx",
                create_test_pptx(),
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".presentationml.presentation"
                ),
            )
        },
    )
    assert response.status_code == 201
    pages = client.get(
        f"/api/documents/{response.json()['document_id']}/pages"
    )
    assert "Briefing" in pages.json()[0]["final_text"]


def test_reject_unknown_extension() -> None:
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "payload.exe",
                b"MZ fake",
                "application/octet-stream",
            )
        },
    )
    assert response.status_code == 400


def test_reject_fake_xlsx_zip() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr(ZipInfo("readme.txt"), "not a workbook")
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "fake.xlsx",
                buffer.getvalue(),
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"
                ),
            )
        },
    )
    assert response.status_code == 400


def test_suspicious_glyph_ratio_and_size_tiers() -> None:
    assert suspicious_glyph_ratio("hello world") == 0.0
    assert suspicious_glyph_ratio("\ufffd\ufffdab") > 0.3
    assert classify_upload_size(1024 * 1024) == "small"
    assert classify_upload_size(20 * 1024 * 1024) == "medium"
    assert prefer_background_processing(20 * 1024 * 1024) is True
