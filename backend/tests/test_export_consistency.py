"""Exports serialize the staged values — never a second extraction.

For each supported document type: the values the Staging Workbook shows
(the grid and the full workbook behind it) equal the JSON export, every
dataset CSV and every Excel cell; the Excel workbook has only the profile's
populated sheets with the source's own headings; long text is preserved."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from uuid import uuid4

import fitz
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.database.session import SessionLocal
from app.main import app
from app.models.document import Document
from app.models.document_page import DocumentPage
from app.services.v3_orchestrator import run_and_persist_v3_extraction
from app.staging import registry
from app.staging.export import _xlsx_value, sheet_table
from app.staging.models import StagingWorkbook
from app.staging.preparation import prepare_staging
from app.staging.profile import SOURCE_SHEET

client = TestClient(app)
ROOT = Path(__file__).resolve().parents[2]
INVOICES = Path(__file__).parent / "fixtures" / "invoices"
REGRESSION = ROOT / "reference" / "regression"


def _process(name: str, content: bytes, mime: str, *, ocr: bool = False) -> str:
    if mime == "application/pdf":
        pdf = fitz.open("pdf", content)
        pdf.set_metadata({"subject": str(uuid4())})
        content = pdf.tobytes()
    else:
        content += f"<!-- {uuid4()} -->".encode()
    upload = client.post("/api/documents/upload", files={"file": (name, content, mime)})
    assert upload.status_code == 201, upload.text
    document_id = upload.json()["document_id"]
    assert client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": ocr}).status_code == 200
    database = SessionLocal()
    try:
        document = database.get(Document, document_id)
        pages = list(database.scalars(select(DocumentPage).where(DocumentPage.document_id == document_id)))
        if mime == "application/pdf":
            run_and_persist_v3_extraction(database=database, document=document, pages=pages)
        prepare_staging(database, document)
    finally:
        database.close()
    return document_id


def _text(value) -> str:
    return "" if value is None else str(value)


def assert_exports_match(document_id: str, profile_id: str) -> dict:
    """The core gate: UI value = JSON value = CSV value = Excel value."""

    full = client.get(f"/api/documents/{document_id}/staging-workbook/export.json").json()
    assert full["profile"]["profile_id"] == profile_id, full["processing_metadata"]
    by_id = {d["dataset_id"]: d for d in full["datasets"]}

    # UI: the grid's values are the workbook's (long text only previewed).
    grid = client.get(f"/api/documents/{document_id}/staging-workbook").json()
    for dataset in grid["datasets"]:
        records = {r["record_id"]: r for r in by_id[dataset["dataset_id"]]["records"]}
        for record in dataset["records"]:
            truncated = set(record.get("truncated_fields") or [])
            for field, cell in record["cells"].items():
                if field not in truncated:
                    assert cell["value"] == records[record["record_id"]]["cells"][field]["value"], (dataset["dataset_id"], field)

    # CSV: each dataset's CSV carries exactly its staged values.
    for dataset in full["datasets"]:
        if not dataset["records"] or dataset["role"] in ("qa",):
            continue
        response = client.get(f"/api/documents/{document_id}/staging-workbook/datasets/{dataset['dataset_id']}.csv")
        assert response.status_code == 200, dataset["dataset_id"]
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        labels = [c["display_label"] for c in dataset["columns"]]
        assert rows[0][: len(labels)] == labels
        for row, record in zip(rows[1:], dataset["records"]):
            expected = [_text(record["cells"][c["canonical_field"]]["value"]) for c in dataset["columns"]]
            assert row[: len(labels)] == expected, dataset["dataset_id"]

    # Excel: the profile's declared sheets, populated only, exact values.
    profile = registry.get_profile(profile_id)
    workbook = StagingWorkbook.model_validate(full)
    model_by_id = {d.dataset_id: d for d in workbook.datasets}
    xlsx = client.get(f"/api/documents/{document_id}/staging-workbook/export.xlsx")
    assert xlsx.status_code == 200
    book = load_workbook(io.BytesIO(xlsx.content))
    expected_sheets = []
    for sheet in profile.export_sheets:
        if sheet.title == SOURCE_SHEET:
            expected_sheets.append(("SOURCE", None))
            continue
        datasets = [model_by_id[i] for i in sheet.dataset_ids if i in model_by_id and any(
            c.value not in (None, "") for r in model_by_id[i].records for c in r.cells.values())]
        if sheet.first_populated:
            datasets = datasets[:1]
        if not datasets:
            continue
        headers, rows = sheet_table(datasets)
        if rows:
            expected_sheets.append(((sheet.title or datasets[0].display_name).upper()[:31], (headers, rows)))
    names = [name for name in book.sheetnames if name != "LONG TEXT"]
    assert names == [name for name, _ in expected_sheets], (names, [n for n, _ in expected_sheets])
    for name, table in expected_sheets:
        ws = book[name]
        values = list(ws.iter_rows(values_only=True))
        assert values[0][0].endswith(name.title()[:5]) or values[0][0]  # title row
        assert values[3][0] is not None  # headings on row 4
        if table is None:
            continue
        headers, rows = table
        assert list(values[3][: len(headers)]) == headers
        assert len(values) - 4 == len(rows), name
        for got, want in zip(values[4:], rows):
            assert list(got[: len(want)]) == [_xlsx_value(v, []) for v in want], name
    return full


def _cell(full: dict, dataset_id: str, label: str):
    dataset = next(d for d in full["datasets"] if d["dataset_id"] == dataset_id)
    for record in dataset["records"]:
        for cell in record["cells"].values():
            if cell["display_label"] == label and cell["value"] is not None:
                return cell["value"]
    return None


def _listed(document_id: str) -> dict:
    database = SessionLocal()
    try:
        name = database.get(Document, document_id).original_filename
    finally:
        database.close()
    found = client.get("/api/documents/search", params={"q": name}).json()["documents"]
    return next(d for d in found if d["document_id"] == document_id)


# --- invoices ---------------------------------------------------------------------------------


def _tesseract() -> bool:
    import shutil

    return shutil.which("tesseract") is not None


@pytest.mark.parametrize(
    "name, ocr",
    [
        ("f1_digital_simple.pdf", False),
        ("f3_multipage_parts.pdf", False),
        ("f4_po_backed_lab.pdf", False),
        ("f5_complex_charges.pdf", False),
        ("f6_html_software.html", False),
        pytest.param("f2_scanned_services.pdf", True, marks=pytest.mark.skipif(not _tesseract(), reason="needs Tesseract")),
    ],
)
def test_invoice_exports_match_the_workbook(name, ocr):
    path = INVOICES / name
    mime = "text/html" if path.suffix == ".html" else "application/pdf"
    document_id = _process(name, path.read_bytes(), mime, ocr=ocr)
    full = assert_exports_match(document_id, "invoice")
    listed = _listed(document_id)
    assert (listed["type_label"], listed["profile_label"]) == ("Invoice", "Invoice")
    # A known value travels unchanged into every format.
    number = _cell(full, "invoice_summary", "Invoice Number")
    if number:
        sheet = load_workbook(io.BytesIO(client.get(f"/api/documents/{document_id}/staging-workbook/export.xlsx").content))["INVOICE SUMMARY"]
        assert number in [row[1] for row in sheet.iter_rows(min_row=5, values_only=True)]
        assert number in client.get(f"/api/documents/{document_id}/staging-workbook/datasets/invoice_summary.csv").text


# --- FAR HTML / FAR clause XML ---------------------------------------------------------------


def test_far_html_exports_match_the_far_view():
    from tests.test_far_part_52 import far_html

    document_id = _process(f"far-{uuid4()}.html", far_html(), "text/html")
    full = client.get(f"/api/documents/{document_id}/staging-workbook/export.json").json()
    records = next(d for d in full["datasets"] if d["dataset_id"] == "far_records")
    labels = [c["display_label"] for c in records["columns"]]
    ui = [[record["cells"][c["canonical_field"]]["value"] for c in records["columns"]] for record in records["records"]]
    base = f"/api/documents/{document_id}/staging-workbook/exports"
    rows = list(csv.reader(io.StringIO(client.get(f"{base}/far_clauses_provisions.csv").content.decode("utf-8-sig"))))
    # The export carries the Part's own schema (Part 52: the clause columns);
    # every exported column holds exactly the value the UI shows.
    exported = rows[0]
    assert exported == [label for label in labels if label in exported] and "FAR Number" in exported
    index = [labels.index(label) for label in exported]
    shown = [[row[i] for i in index] for row in ui]
    assert rows[1:] == [[_text(v) for v in row] for row in shown]
    sheet = list(load_workbook(io.BytesIO(client.get(f"{base}/far_workbook.xlsx").content))["FAR CLAUSES & PROVISIONS"].iter_rows(values_only=True))
    assert list(sheet[3]) == exported
    assert [list(row) for row in sheet[4:]] == shown
    listed = _listed(document_id)
    assert (listed["type_label"], listed["profile_label"]) == ("FAR Part 52", "FAR Regulation")


def test_far_clause_xml_exports_match_the_workbook():
    from tests.test_xml_document import CLAUSE_XML

    long_text = "(a) " + "The Contractor shall comply. " * 1300  # beyond an Excel cell
    xml = CLAUSE_XML.replace("<p>Definitions (Jun 2020)</p>", f"<p>Definitions (Jun 2020)</p><p>{long_text}</p>")
    document_id = _process("FAR_Part_52_OKCXMLIMPDFN_T.xml", xml.encode(), "application/xml")
    full = assert_exports_match(document_id, "xml_document")
    clauses = next(d for d in full["datasets"] if d["dataset_id"] == "xml_records_1")
    assert [c["display_label"] for c in clauses["columns"]][-1] == "Source XML"
    # Long text: complete in JSON and CSV, continued in LONG TEXT in Excel.
    text = clauses["records"][1]["cells"][next(c["canonical_field"] for c in clauses["columns"] if c["display_label"] == "Text")]["value"]
    assert long_text.strip() in text
    assert long_text.strip() in client.get(f"/api/documents/{document_id}/staging-workbook/datasets/xml_records_1.csv").text
    book = load_workbook(io.BytesIO(client.get(f"/api/documents/{document_id}/staging-workbook/export.xlsx").content))
    assert "LONG TEXT" in book.sheetnames
    listed = _listed(document_id)
    assert (listed["type_label"], listed["profile_label"]) == ("FAR Clause XML", "XML Document")


# --- contracts --------------------------------------------------------------------------------


def _contract(path: Path, expected_type: str) -> None:
    document_id = _process(path.name, path.read_bytes(), "application/pdf")
    full = assert_exports_match(document_id, "contract_v3")
    book = load_workbook(io.BytesIO(client.get(f"/api/documents/{document_id}/staging-workbook/export.xlsx").content))
    # Combined sheets, never a dozen empty legacy ones.
    assert set(book.sheetnames) - {"LONG TEXT"} <= {
        "CONTRACT SUMMARY", "CONTRACT DATA", "LINE ITEMS", "CLAUSES & PROVISIONS", "OTHER INFORMATION", "SOURCE",
    }
    data = next(d for d in full["datasets"] if d["dataset_id"] == "contract_data")
    assert book["CONTRACT DATA"].max_row - 4 == len(data["records"])
    listed = _listed(document_id)
    assert (listed["type_label"], listed["profile_label"]) == (expected_type, "Government Contract")


@pytest.mark.skipif(not (REGRESSION / "SF1442 Award_N4019223D2803.pdf").exists(), reason="regression PDF not available")
def test_sf1442_exports_match_the_workbook():
    _contract(REGRESSION / "SF1442 Award_N4019223D2803.pdf", "SF 1442 Contract Award")


@pytest.mark.skipif(not (REGRESSION / "Contract_47QRCA25DSF07 (2).pdf").exists(), reason="regression PDF not available")
def test_47qrc_exports_match_the_workbook():
    _contract(REGRESSION / "Contract_47QRCA25DSF07 (2).pdf", "Government Contract")


@pytest.mark.skipif(not (REGRESSION / "Contract - FA300224C0008 (2).pdf").exists(), reason="regression PDF not available")
def test_fa30_exports_match_the_workbook():
    _contract(REGRESSION / "Contract - FA300224C0008 (2).pdf", "Government Contract")


# --- transcript -------------------------------------------------------------------------------


def test_transcript_exports_match_the_workbook():
    from tests.test_academic_transcript import _pdf, _university

    document_id = _process(f"transcript-{uuid4()}.pdf", _pdf(_university), "application/pdf")
    assert_exports_match(document_id, "academic_transcript")
    listed = _listed(document_id)
    assert (listed["type_label"], listed["profile_label"]) == ("Academic Transcript", "Academic Transcript")
