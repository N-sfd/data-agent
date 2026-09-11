"""Ground-truth fixtures for Universal Document Ingestion OCR matrix."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches as PptxInches

INVOICE_GROUND_TRUTH = {
    "Invoice Number": "INV-2026-8914",
    "Invoice Date": "09/05/2026",
    "Purchase Order": "PO-78452",
    "TOTAL": "$30,500.00",
    "Payment Terms": "Net 30",
}

CONTRACT_GROUND_TRUTH = {
    "Contract Number": "CA-2026-0047",
    "Effective Date": "September 1, 2026",
    "Expiration Date": "August 31, 2027",
    "Vendor": "Consult America, Inc.",
    "Total Contract Value": "$485,750.00",
}

# Format families → whether OCR is expected for the primary content path.
OCR_EXPECTATIONS = {
    "raster_image": True,  # JPG/PNG/BMP/WEBP/TIFF
    "image_only_pdf": True,
    "mixed_pdf_native_page": False,
    "mixed_pdf_scanned_page": True,
    "docx_text": False,
    "docx_with_scanned_image": True,  # OCR on embedded image content
    "pptx_with_scanned_slide": True,
    "xlsx": False,
    "txt": False,
    "csv": False,
    "html": False,
    "rtf": False,
}


def _font(size: int = 28):
    try:
        return ImageFont.truetype("arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def render_ground_truth_image(
    title: str,
    fields: dict[str, str],
    *,
    size: tuple[int, int] = (1200, 1600),
) -> bytes:
    """Render labeled ground-truth fields onto a high-res scan-like image."""

    image = Image.new("RGB", size, color="white")
    draw = ImageDraw.Draw(image)
    title_font = _font(36)
    body_font = _font(28)
    draw.text((64, 64), title, fill="black", font=title_font)
    y = 160
    for label, value in fields.items():
        draw.text((64, y), f"{label}: {value}", fill="black", font=body_font)
        y += 70

    buffer = BytesIO()
    image.save(buffer, format="PNG", dpi=(200, 200))
    return buffer.getvalue()


def image_only_pdf_from_png(png_bytes: bytes) -> bytes:
    pdf = fitz.open()
    page = pdf.new_page(width=612, height=792)
    page.insert_image(page.rect, stream=png_bytes)
    data = pdf.tobytes()
    pdf.close()
    return data


def mixed_native_and_scanned_pdf(
    *,
    native_text: str,
    scanned_png: bytes,
) -> bytes:
    pdf = fitz.open()
    native = pdf.new_page()
    native.insert_text((72, 72), native_text)
    scanned = pdf.new_page()
    scanned.insert_image(scanned.rect, stream=scanned_png)
    data = pdf.tobytes()
    pdf.close()
    return data


def native_text_pdf(text: str) -> bytes:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), text)
    data = pdf.tobytes()
    pdf.close()
    return data


def docx_with_embedded_scan(png_bytes: bytes, *, caption: str) -> bytes:
    document = DocxDocument()
    document.add_paragraph(caption)
    document.add_picture(BytesIO(png_bytes), width=Inches(6.0))
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_with_scanned_slide(png_bytes: bytes, *, title: str) -> bytes:
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(blank)
    slide.shapes.add_picture(
        BytesIO(png_bytes),
        PptxInches(0.5),
        PptxInches(0.5),
        width=PptxInches(9),
    )
    box = slide.shapes.add_textbox(
        PptxInches(0.5), PptxInches(7.0), PptxInches(9), PptxInches(0.4)
    )
    box.text_frame.paragraphs[0].text = title
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def xlsx_native_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet["A1"] = "Item"
    sheet["B1"] = "Amount"
    sheet["A2"] = "Widget"
    sheet["B2"] = 30500
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def write_fixture(path: Path, data: bytes) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path
