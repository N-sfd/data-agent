"""CSV / Complete-Excel builders for the canonical V3 document
(docs/v3-schema-manifest.md). Every builder serializes the already-
assembled `NormalizedV3Document` — no re-extraction, no independent
re-derivation, per §17 of the P0 brief ("Export must serialize the
persisted canonical V3 data").
"""

from __future__ import annotations

import csv
import io
from typing import Any

from app.schemas.v3_document import NormalizedV3Document

# Static — reproduces the ground-truth workbook's README principles
# verbatim (docs/v3-schema-manifest.md §12). Not extracted business data.
README_ROWS: list[tuple[str, str]] = [
    (
        "GROUND TRUTH",
        "Only source-supported values are populated; missing values are not guessed.",
    ),
    (
        "PROVENANCE",
        "Every substantive record includes source page and exact/near-exact evidence.",
    ),
    (
        "CLINS",
        "Repeating line items are isolated from ordinary contract metadata.",
    ),
    (
        "FUNDING",
        "Basic-contract funding facts are separated from task-order-level funding requirements.",
    ),
    (
        "CLAUSES",
        "Contract clauses are separated from incidental FAR references.",
    ),
    (
        "QA",
        "Ambiguous/missing source material is explicitly flagged for review.",
    ),
]

_CLIN_HEADERS = [
    "CLIN", "Option/Base", "Description", "Pricing Type", "Max Quantity", "Unit",
    "Unit Price", "Max Amount", "Status", "FOB", "Purchase Request", "PSC",
    "POP Start", "POP End", "Ship To", "DODAAC", "Source Page", "Evidence", "QA Status",
]
_FUNDING_HEADERS = [
    "Funding Level", "CLIN", "Funding Status", "Amount", "Accounting / Appropriation",
    "Purchase Request", "Source Page", "Evidence", "QA Status",
]
_PERFORMANCE_HEADERS = [
    "Record Type", "CLIN", "Start", "End / Timing", "Location / Destination",
    "Requirement", "Source Page", "Evidence", "QA Status",
]
_ATTACHMENT_HEADERS = [
    "Attachment / Reference", "Title / Description", "Included in Portfolio?",
    "Source Page", "Evidence", "QA Status",
]
_CLAUSE_HEADERS = [
    "Regulation", "Clause Number", "Clause Title", "Alternate / Deviation",
    "Effective Date", "Incorporation Type", "Source Page", "Evidence", "QA Status",
]
_FAR_REFERENCE_HEADERS = [
    "FAR Reference", "Reference Type", "Subject / Context", "Source Page",
    "Evidence", "Contract Clause?", "QA Status",
]
_SOURCE_DOCUMENT_HEADERS = ["Source Document", "Role", "Pages", "Extraction Status"]
_QA_REVIEW_HEADERS = ["QA Check", "Result", "Details", "Action"]
_ALL_FIELDS_HEADERS = [
    "Category", "Normalized Field", "Value", "Source File", "Source Page",
    "Evidence", "Extraction Method", "QA Status",
]
_CONTRACT_SUMMARY_HEADERS = [
    "Contract Number", "Solicitation / RFP", "Contract Vehicle", "Agency / Office",
    "Contractor", "Award Date", "Ceiling / Max Aggregate", "Minimum Guarantee",
    "Base Period", "Options", "Max Duration", "Task Order Range", "NAICS",
    "Size Standard", "Source File", "Source Page", "Evidence", "QA Status",
]


def _write_csv(headers: list[str], rows: list[list[Any]]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if value is None else value for value in row])
    return buffer.getvalue()


def _clin_rows(doc: NormalizedV3Document) -> list[list[Any]]:
    return [
        [
            r.clin, r.option_base, r.description, r.pricing_type, r.max_quantity, r.unit,
            r.unit_price, r.max_amount, r.status, r.fob, r.purchase_request, r.psc,
            r.pop_start, r.pop_end, r.ship_to, r.dodaac, r.source_page, r.evidence, r.qa_status,
        ]
        for r in doc.clins
    ]


def _funding_rows(doc: NormalizedV3Document) -> list[list[Any]]:
    return [
        [
            r.funding_level, r.clin, r.funding_status, r.amount, r.accounting_appropriation,
            r.purchase_request, r.source_page, r.evidence, r.qa_status,
        ]
        for r in doc.funding
    ]


def _performance_rows(doc: NormalizedV3Document) -> list[list[Any]]:
    return [
        [
            r.record_type, r.clin, r.start, r.end_timing, r.location_destination,
            r.requirement, r.source_page, r.evidence, r.qa_status,
        ]
        for r in doc.performance_delivery
    ]


def _attachment_rows(doc: NormalizedV3Document) -> list[list[Any]]:
    return [
        [
            r.attachment_reference, r.title_description, r.included_in_portfolio,
            r.source_page, r.evidence, r.qa_status,
        ]
        for r in doc.attachments
    ]


def _clause_rows(rows) -> list[list[Any]]:
    return [
        [
            r.regulation, r.clause_number, r.clause_title, r.alternate_deviation,
            r.effective_date, r.incorporation_type, r.source_page, r.evidence, r.qa_status,
        ]
        for r in rows
    ]


def _far_reference_rows(doc: NormalizedV3Document) -> list[list[Any]]:
    return [
        [
            r.far_reference, r.reference_type, r.subject_context, r.source_page,
            r.evidence, r.contract_clause, r.qa_status,
        ]
        for r in doc.far_references
    ]


def _source_document_rows(doc: NormalizedV3Document) -> list[list[Any]]:
    return [[r.source_document, r.role, r.pages, r.extraction_status] for r in doc.source_documents]


def _qa_review_rows(doc: NormalizedV3Document) -> list[list[Any]]:
    return [[r.qa_check, r.result, r.details, r.action] for r in doc.qa_review]


def _all_fields_rows(doc: NormalizedV3Document) -> list[list[Any]]:
    return [
        [
            r.category, r.normalized_field, r.value, r.source_file, r.source_page,
            r.evidence, r.extraction_method, r.qa_status,
        ]
        for r in doc.all_fields
    ]


def _contract_summary_row(doc: NormalizedV3Document) -> list[Any]:
    s = doc.contract_summary
    if s is None:
        return [""] * len(_CONTRACT_SUMMARY_HEADERS)
    return [
        s.contract_number, s.solicitation_rfp, s.contract_vehicle, s.agency_office,
        s.contractor, s.award_date, s.ceiling_max_aggregate, s.minimum_guarantee,
        s.base_period, s.options, s.max_duration, s.task_order_range, s.naics,
        s.size_standard, s.source_file, s.source_page, s.evidence, s.qa_status,
    ]


_DATASET_BUILDERS: dict[str, tuple[list[str], Any]] = {
    "all-fields": (_ALL_FIELDS_HEADERS, _all_fields_rows),
    "clins": (_CLIN_HEADERS, _clin_rows),
    "funding": (_FUNDING_HEADERS, _funding_rows),
    "performance-delivery": (_PERFORMANCE_HEADERS, _performance_rows),
    "attachments": (_ATTACHMENT_HEADERS, _attachment_rows),
    "clauses": (_CLAUSE_HEADERS, lambda doc: _clause_rows(doc.clauses)),
    "far-references": (_FAR_REFERENCE_HEADERS, _far_reference_rows),
    "dfars": (_CLAUSE_HEADERS, lambda doc: _clause_rows(doc.dfars)),
    "source-documents": (_SOURCE_DOCUMENT_HEADERS, _source_document_rows),
    "qa-review": (_QA_REVIEW_HEADERS, _qa_review_rows),
}

DATASET_NAMES = tuple(_DATASET_BUILDERS.keys()) + ("contract-summary",)


def build_dataset_csv(doc: NormalizedV3Document, dataset: str) -> str:
    if dataset == "contract-summary":
        return _write_csv(_CONTRACT_SUMMARY_HEADERS, [_contract_summary_row(doc)])
    entry = _DATASET_BUILDERS.get(dataset)
    if entry is None:
        raise ValueError(f"Unknown V3 dataset: {dataset!r}")
    headers, row_builder = entry
    return _write_csv(headers, row_builder(doc))


def build_v3_xlsx(doc: NormalizedV3Document) -> bytes:
    """12 sheets, exact ground-truth order/names/headers
    (docs/v3-schema-manifest.md)."""

    from openpyxl import Workbook

    workbook = Workbook()
    sheet_order = [
        ("All Fields", _ALL_FIELDS_HEADERS, _all_fields_rows(doc)),
        ("CLINs", _CLIN_HEADERS, _clin_rows(doc)),
        ("Funding", _FUNDING_HEADERS, _funding_rows(doc)),
        ("Performance Delivery", _PERFORMANCE_HEADERS, _performance_rows(doc)),
        ("Attachments", _ATTACHMENT_HEADERS, _attachment_rows(doc)),
        ("Clauses", _CLAUSE_HEADERS, _clause_rows(doc.clauses)),
        ("FAR References", _FAR_REFERENCE_HEADERS, _far_reference_rows(doc)),
        ("DFARS", _CLAUSE_HEADERS, _clause_rows(doc.dfars)),
        ("Source Documents", _SOURCE_DOCUMENT_HEADERS, _source_document_rows(doc)),
        ("QA Review", _QA_REVIEW_HEADERS, _qa_review_rows(doc)),
        ("Contract Summary", _CONTRACT_SUMMARY_HEADERS, [_contract_summary_row(doc)]),
    ]

    first = True
    for title, headers, rows in sheet_order:
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = title
        sheet.append(headers)
        for row in rows:
            sheet.append(["" if value is None else value for value in row])

    readme_sheet = workbook.create_sheet("README")
    readme_sheet.append(["Principle", "Rule"])
    for principle, rule in README_ROWS:
        readme_sheet.append([principle, rule])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
