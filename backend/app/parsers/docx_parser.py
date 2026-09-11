from docx import Document

from app.services.embedded_image_ocr import ocr_image_bytes


def parse_docx(file_path: str) -> dict:
    document = Document(file_path)

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    tables = []

    for table in document.tables:
        rows = []

        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])

        tables.append(rows)

    ocr_snippets: list[str] = []
    for rel in document.part.rels.values():
        rel_type = getattr(rel, "reltype", "") or ""
        if "image" not in rel_type:
            continue
        try:
            blob = rel.target_part.blob
        except Exception:
            continue
        text = ocr_image_bytes(blob)
        if text:
            ocr_snippets.append(text)
            paragraphs.append(text)

    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "ocr_used": bool(ocr_snippets),
        "ocr_snippets": ocr_snippets,
    }
