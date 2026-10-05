"""FAR Part 52 exports, built from the persisted canonical records — never
from a reference workbook:

- the FAR workbook: FAR CLAUSES & PROVISIONS, ORACLE OUTPUT, SOURCE (and
  LONG TEXT only when a value exceeds an Excel cell), plus one CSV each;
- the technical transformation workbook (01–08 + 99) and its CSVs;
- the canonical model as JSON."""

from __future__ import annotations

import csv
import io
import json

from app.far import regulation, views
from app.far.oracle_map import NOTICE, TITLE
from app.models.document_far_record import DocumentFarRecord

SHEETS = {
    "01_FAR_STAGING": "FAR Part 52 — Grouped Professional Staging",
    "02_ORACLE_MAPPING": "FAR Part 52 — Oracle Mapping",
    "03_VALIDATION": "FAR Part 52 — Validation Checks",
    "99_LONG_TEXT": "FAR Part 52 — Long Text",
    "04_STRUCTURED_FAR": "FAR Part 52 — Structured Clause / Provision Extraction",
    "05_EXTRACTION_GUIDE": "FAR Part 52 — Extraction Rules",
    "06_CANONICAL_MODEL": "Canonical FAR Transformation Model",
    "07_ORACLE_OUTPUT_MAP": TITLE,
    "08_BUSINESS_EXPORT": "FAR Part 52 — Business Export",
}
SUBTITLES = {
    "02_ORACLE_MAPPING": "Each FAR record with its Oracle clause-library target values. Blank means no approved "
    "rule or target template supports a value yet.",
    "03_VALIDATION": "Checks that flag a record for review; they never change a value.",
    "99_LONG_TEXT": "Text longer than an Excel cell, in full, as consecutive chunks of each FAR record.",
    "05_EXTRACTION_GUIDE": "How each field is extracted and formatted from the FAR HTML source.",
    "01_FAR_STAGING": "One FAR record per row, structural alternates on their own rows after their basic clause. "
    "Text is kept verbatim (never rewritten). Paragraph/Subparagraph remain the last fields.",
    "04_STRUCTURED_FAR": "One FAR record per row. Dates, prescriptions, structural alternates and embedded FAR "
    "references are separated; headings and text are kept intact.",
    "06_CANONICAL_MODEL": "System-neutral canonical layer between FAR extraction and downstream Oracle output. "
    "Alternates are separate records linked by Basic Clause Key; structural and reserved records are kept with "
    "Load Eligible = NO.",
    "07_ORACLE_OUTPUT_MAP": NOTICE,
    "08_BUSINESS_EXPORT": (
        "Every FAR record in the clause-library layout, with Type and the text split into Provision Text "
        "and Clause Text. Library defaults apply to load-eligible records; Source XML stays blank until "
        "the import file naming is supplied. Text is complete in CSV and JSON; Excel cells follow the "
        "existing long-text limit and 99_LONG_TEXT."
    ),
}


def sheet_rows(sheet: str, rows: list[DocumentFarRecord], summary: dict | None) -> tuple[tuple[str, ...], list[dict]]:
    if sheet == "01_FAR_STAGING":
        return views.STAGING_COLUMNS, [views.staging_row(r) for r in rows]
    if sheet == "02_ORACLE_MAPPING":
        return views.ORACLE_MAPPING_COLUMNS, [views.oracle_mapping_row(r) for r in rows]
    if sheet == "03_VALIDATION":
        return views.VALIDATION_COLUMNS, views.validation_rows(rows, summary)
    if sheet == "99_LONG_TEXT":
        return views.LONG_TEXT_COLUMNS, views.long_text_rows(rows)
    if sheet == "04_STRUCTURED_FAR":
        return views.STRUCTURED_COLUMNS, [views.structured_row(r) for r in rows]
    if sheet == "05_EXTRACTION_GUIDE":
        columns = ("Field", "Extraction / Formatting Rule", "Data Agent Use")
        return columns, [dict(zip(columns, row)) for row in views.EXTRACTION_GUIDE]
    if sheet == "06_CANONICAL_MODEL":
        return views.CANONICAL_COLUMNS, [views.canonical_row(r) for r in rows]
    if sheet == "07_ORACLE_OUTPUT_MAP":
        return views.OUTPUT_MAP_COLUMNS, views.output_map_rows()
    if sheet == "08_BUSINESS_EXPORT":
        return views.BUSINESS_COLUMNS, [views.business_row(r) for r in rows]
    raise KeyError(sheet)


# Paragraph-length columns get a wide, wrapped column.
WIDE_COLUMNS = {
    "Actual Section Text",
    "Source Text",
    "Text",
    "Provision Text",
    "Clause Text",
    "Section Text",
    "Description",
    "Transformation Rule",
    "Notes",
    "Extraction / Formatting Rule",
    "Data Agent Use",
}


def build_far_transformation_workbook(rows: list[DocumentFarRecord], summary: dict | None) -> bytes:
    from openpyxl import Workbook

    from app.staging.xlsx_style import write_table_sheet

    wb = Workbook()
    wb.remove(wb.active)
    for sheet, title in SHEETS.items():
        columns, records = sheet_rows(sheet, rows, summary)
        values = []
        for record in records:
            row = []
            for column in columns:
                value = record.get(column)
                # A whole long text lives in 99_LONG_TEXT; the cell
                # carries its first chunk (Excel's cell limit).
                if isinstance(value, str) and len(value) > views.LONG_TEXT_CHUNK and sheet != "99_LONG_TEXT":
                    value = value[: views.LONG_TEXT_CHUNK]
                row.append(value)
            values.append(row)
        write_table_sheet(
            wb.create_sheet(sheet),
            title=title,
            subtitle=SUBTITLES.get(sheet),
            columns=list(columns),
            rows=values,
            wide_columns=WIDE_COLUMNS,
        )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_far_csv(sheet: str, rows: list[DocumentFarRecord], summary: dict | None) -> str:
    columns, records = sheet_rows(sheet, rows, summary)
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(columns), extrasaction="ignore")
    writer.writeheader()
    writer.writerows(records)
    return buffer.getvalue()


def build_far_json(rows: list[DocumentFarRecord], summary: dict | None, document: dict) -> str:
    parents = views.far_parents(rows)
    payload = {
        "document": document,
        "profile": "far_part_52@1",
        "summary": views.overview(rows, summary),
        "canonical_model": [
            {
                **views.canonical_row(r),
                "Official Heading": r.official_heading,
                "Clause / Provision": r.clause_type,
                "Prescription": r.prescription,
                "Alternate Heading": r.alternate_heading,
                "Alternate Instruction": r.alternate_instruction,
                "Section": r.subpart,
                "Subsection": r.subsection,
                "Embedded FAR References": list(r.embedded_references or []),
                "Source Locator": r.provenance_json,
            }
            for r in rows
        ],
        "validation": views.validation_rows(rows, summary),
        "oracle_output_mapping": {"title": TITLE, "notice": NOTICE, "rows": views.output_map_rows()},
        "business_export": [views.business_row(r) for r in rows],
        "far_clauses_provisions": far_sheet_rows("FAR CLAUSES & PROVISIONS", rows, summary, document.get("filename"))[1],
        # Explicit citations per record, typed (FAR, CFR, USC, Public Law,
        # Executive Order, Form) — the structured form of Cross References.
        "references": [
            {
                "far_number": r.far_number,
                "clause_key": r.clause_key,
                "references": [
                    {"type": ref.kind, "reference": ref.text}
                    for ref in regulation.typed_references(r.source_text, r.far_number)
                ],
            }
            for r in rows
        ],
        "oracle_output": [views.oracle_output_row(r) for r in rows],
    }
    return json.dumps(payload, ensure_ascii=False, indent=1)


CSV_SHEETS = (
    "01_FAR_STAGING",
    "02_ORACLE_MAPPING",
    "03_VALIDATION",
    "04_STRUCTURED_FAR",
    "06_CANONICAL_MODEL",
    "07_ORACLE_OUTPUT_MAP",
    "08_BUSINESS_EXPORT",
)


# --- the FAR workbook -------------------------------------------------------------------------

# Sheet → (title, subtitle). The first sheet is FAR semantics only; the
# Oracle target shape lives on its own sheet.
FAR_SHEETS = {
    "FAR CLAUSES & PROVISIONS": (
        "FAR Part 52 — Clauses & Provisions",
        "Every FAR record in source order. An alternate follows its basic record and keeps its FAR Number. "
        "A provision's body is in Provision Text, a clause's in Clause Text, a regulation section's in "
        "Section Text — one per record. Text is verbatim.",
    ),
    "ORACLE OUTPUT": (
        "FAR Part 52 — Oracle Output",
        "The FAR records shaped for the Oracle Fusion clause library. Library values apply to load-eligible "
        "records only; structural, reserved and alternate records carry none.",
    ),
    "SOURCE": (
        "FAR Part 52 — Source",
        "The source document, what it contains, and how every value traces back to it.",
    ),
}
FAR_CSV = {
    "far_clauses_provisions": "FAR CLAUSES & PROVISIONS",
    "oracle_output": "ORACLE OUTPUT",
    "source": "SOURCE",
}
_CELL_LIMIT = 32_767


def part_label(rows: list[DocumentFarRecord]) -> str:
    part = views.document_part(rows)
    return f"FAR Part {part}" if part else "FAR"


def sheet_heading(sheet: str, rows: list[DocumentFarRecord]) -> tuple[str, str, str]:
    """(sheet name, title, subtitle) for this Part: Part 52 keeps "Clauses &
    Provisions"; any other Part's record table is its FAR Data."""

    label = part_label(rows)
    title, subtitle = FAR_SHEETS[sheet]
    title = title.replace("FAR Part 52", label)
    if sheet == "FAR CLAUSES & PROVISIONS" and views.document_part(rows) != "52":
        return (
            "FAR DATA",
            f"{label} — FAR Data",
            "Every numbered section and subsection of the Part in source order, under its Part and Subpart. "
            "Record Type is the record's regulatory function; Prescription / Usage and Prescription Reference "
            "name the provisions, clauses or forms it prescribes. Paragraphs (a), (1), (i) stay inside the "
            "regulatory text, which is verbatim.",
        )
    return sheet, title, subtitle


def far_sheet_rows(
    sheet: str, rows: list[DocumentFarRecord], summary: dict | None, filename: str | None
) -> tuple[tuple[str, ...], list[dict]]:
    if sheet == "FAR CLAUSES & PROVISIONS":
        parents = views.far_parents(rows)
        part_title = regulation.part_title((summary or {}).get("part_heading"))
        columns = views.record_columns(views.document_part(rows))
        records = [views.far_record_row(r, parents, filename, part_title) for r in rows]
        return tuple(header for header, _ in columns), [{header: record.get(key) for header, key in columns} for record in records]
    if sheet == "ORACLE OUTPUT":
        return views.ORACLE_OUTPUT_COLUMNS, [views.oracle_output_row(r) for r in rows]
    if sheet == "SOURCE":
        return ("Item", "Detail"), views.source_rows(rows, summary, filename)
    raise KeyError(sheet)


def build_far_workbook(rows: list[DocumentFarRecord], summary: dict | None, filename: str | None = None) -> bytes:
    from openpyxl import Workbook

    from app.staging.xlsx_style import write_table_sheet

    wb = Workbook()
    wb.remove(wb.active)
    long_text: list[list] = []
    for sheet in FAR_SHEETS:
        sheet_name, title, subtitle = sheet_heading(sheet, rows)
        columns, records = far_sheet_rows(sheet, rows, summary, filename)
        values = []
        for record in records:
            row = []
            for column in columns:
                value = record.get(column)
                if isinstance(value, str) and len(value) > _CELL_LIMIT:
                    # The whole text goes to LONG TEXT, in parts; the cell
                    # keeps its start and says where the rest is.
                    parts = [value[i : i + views.LONG_TEXT_CHUNK] for i in range(0, len(value), views.LONG_TEXT_CHUNK)]
                    number = record.get("FAR Number") or record.get("Number")
                    label = " ".join(p for p in (number, record.get("Alternate")) if p)
                    long_text += [[sheet_name, label, column, f"{n} of {len(parts)}", part] for n, part in enumerate(parts, 1)]
                    value = value[: views.LONG_TEXT_CHUNK] + f" … [continued in LONG TEXT: {len(value):,} characters]"
                row.append(value)
            values.append(row)
        source = f" Source: {filename}." if filename else ""
        write_table_sheet(
            wb.create_sheet(sheet_name),
            title=title,
            subtitle=subtitle + source,
            columns=list(columns),
            rows=values,
            wide_columns=WIDE_COLUMNS | {"Prescription / Usage", "Source Reference", "Detail", "Title"},
        )
    if long_text:
        write_table_sheet(
            wb.create_sheet("LONG TEXT"),
            title=f"{part_label(rows)} — Long Text",
            subtitle="Values longer than an Excel cell holds, in full, as consecutive parts.",
            columns=["Sheet", "Record", "Field", "Part", "Text"],
            rows=long_text,
            wide_columns={"Text"},
        )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_far_sheet_csv(sheet: str, rows: list[DocumentFarRecord], summary: dict | None, filename: str | None) -> str:
    columns, records = far_sheet_rows(sheet, rows, summary, filename)
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(columns), extrasaction="ignore")
    writer.writeheader()
    writer.writerows(records)
    return buffer.getvalue()
