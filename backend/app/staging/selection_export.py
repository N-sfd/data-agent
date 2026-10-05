"""Selected export: the cells a reviewer picked in a staging grid, as Excel,
CSV or JSON.

The browser sends a manifest — which records and which fields, by their
stable ids — never values. Values come from the document's staging
workbook, built from its persisted records exactly as Export All builds it
(no OCR, extraction, classification or AI runs). Every record id must
belong to the dataset and every field must be one of its business columns.

The output keeps the grid's table shape: the selected fields as columns in
the dataset's own order, the selected records as rows in source order. A
record's identifier column (e.g. FAR Number) is always carried so a sparse
selection still reads; a cell that wasn't selected stays empty.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from typing import Literal

from pydantic import BaseModel, Field

from app.staging.export import _sheet_title, _write_long_text, _xlsx_value, find_dataset
from app.staging.models import StagingColumn, StagingDataset, StagingWorkbook

ExportFormat = Literal["xlsx", "csv", "json"]

# Keys a column must never carry to be exported: internal ids, geometry,
# extraction internals and raw provenance. Business columns (including a
# Source Reference column) are never named like these.
_TECHNICAL_KEYS = {
    "record_id",
    "field_id",
    "database_id",
    "bbox",
    "source_bbox",
    "confidence",
    "extraction_method",
    "ocr_method",
    "provenance",
    "provenance_json",
    "dom_path",
}
_TECHNICAL_SUFFIXES = ("_bbox", "_confidence", "_provenance")
_MAX_CELLS = 2_000_000


class SelectedRecord(BaseModel):
    record_id: str
    fields: list[str] = Field(min_length=1)


class SelectionManifest(BaseModel):
    format: ExportFormat
    selection: list[SelectedRecord] = Field(min_length=1)


class SelectionError(ValueError):
    """The manifest names a record or field the dataset doesn't export."""


@dataclass
class SelectedTable:
    heading: str
    subtitle: str
    columns: list[StagingColumn]
    # Per record: (record_id, {canonical_field: value}) for the exported cells.
    rows: list[tuple[str, dict[str, object]]]
    identifier: str | None = None


def is_exportable(column: StagingColumn) -> bool:
    key = column.key.lower()
    return key not in _TECHNICAL_KEYS and not key.endswith(_TECHNICAL_SUFFIXES)


def exportable_columns(dataset: StagingDataset) -> list[StagingColumn]:
    return [column for column in dataset.columns if is_exportable(column)]


def identifier_field(dataset: StagingDataset, columns: list[StagingColumn]) -> str | None:
    """The column that names a record: the dataset's first identity field,
    else its first grid column, else its first column (none for a
    single-record dataset)."""

    if dataset.cardinality == "single":
        return None  # one record: nothing to tell apart, export exactly what was chosen
    exportable = {column.canonical_field for column in columns}
    for field in [*dataset.identity_fields, *dataset.grid_fields]:
        if field in exportable:
            return field
    return columns[0].canonical_field if columns else None


def _heading(workbook: StagingWorkbook, dataset: StagingDataset) -> str:
    """"FAR Part 49 — Selected Records" for a FAR Part; otherwise the
    dataset's own name."""

    source = find_dataset(workbook, "far_source")
    if source is not None:
        for record in source.records:
            if record.cells.get("far.source.item") and record.cells["far.source.item"].value == "Part":
                part = re.search(r"\bPart\s+(\d+)", str(record.cells["far.source.detail"].value or ""))
                if part:
                    return f"FAR Part {part.group(1)} — Selected Records"
    return f"{dataset.display_name} — Selected Records"


def select(workbook: StagingWorkbook, dataset_id: str, manifest: SelectionManifest) -> SelectedTable:
    dataset = find_dataset(workbook, dataset_id)
    if dataset is None:
        raise SelectionError(f"Unknown dataset {dataset_id!r}.")
    columns = exportable_columns(dataset)
    allowed = {column.canonical_field for column in columns}

    wanted: dict[str, set[str]] = {}
    for item in manifest.selection:
        wanted.setdefault(item.record_id, set()).update(item.fields)
    if sum(len(fields) for fields in wanted.values()) > _MAX_CELLS:
        raise SelectionError("The selection is too large to export.")

    unknown_fields = sorted({f for fields in wanted.values() for f in fields} - allowed)
    if unknown_fields:
        raise SelectionError(f"Fields not exportable from {dataset_id!r}: {', '.join(unknown_fields[:10])}.")
    records = {record.record_id: record for record in dataset.records}
    unknown_records = [record_id for record_id in wanted if record_id not in records]
    if unknown_records:
        raise SelectionError(
            f"{len(unknown_records)} record(s) are not in this document's {dataset.display_name}: "
            f"{', '.join(unknown_records[:5])}."
        )

    identifier = identifier_field(dataset, columns)
    chosen = set().union(*wanted.values())
    shown = [c for c in columns if c.canonical_field in chosen or c.canonical_field == identifier]
    rows: list[tuple[str, dict[str, object]]] = []
    for record in dataset.records:  # source order, never selection order
        fields = wanted.get(record.record_id)
        if not fields:
            continue
        exported = fields | ({identifier} if identifier else set())
        rows.append(
            (
                record.record_id,
                {
                    column.canonical_field: record.cells[column.canonical_field].value
                    for column in shown
                    if column.canonical_field in exported and column.canonical_field in record.cells
                },
            )
        )

    noun = "record" if len(rows) == 1 else "records"
    field_count = len(shown)
    subtitle = (
        f"Extracted from {workbook.document_filename} · {dataset.display_name} · "
        f"{len(rows)} selected {noun}, {field_count} field{'s' if field_count != 1 else ''}."
    )
    return SelectedTable(_heading(workbook, dataset), subtitle, shown, rows, identifier)


def _filename(workbook: StagingWorkbook, dataset_id: str, extension: str) -> str:
    stem = (workbook.document_filename or "export").rsplit(".", 1)[0]
    return f"{stem}_{dataset_id}_selected.{extension}"


def build_selected_csv(table: SelectedTable) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([column.display_label for column in table.columns])
    for _record_id, values in table.rows:
        writer.writerow([values.get(column.canonical_field) for column in table.columns])
    return buffer.getvalue()


def build_selected_json(table: SelectedTable, workbook: StagingWorkbook, dataset: StagingDataset) -> str:
    """Semantic field names (a column's key, e.g. "section_text"); a record
    carries only its exported cells."""

    payload = {
        "document": workbook.document_filename,
        "dataset": dataset.display_name,
        "dataset_id": dataset.dataset_id,
        "fields": [{"key": c.key, "label": c.display_label, "canonical_field": c.canonical_field} for c in table.columns],
        # The column every record carries so a sparse selection still reads.
        "identifier": next((c.key for c in table.columns if c.canonical_field == table.identifier), None),
        "records": [
            {c.key: values[c.canonical_field] for c in table.columns if c.canonical_field in values}
            for _record_id, values in table.rows
        ],
    }
    return json.dumps(payload, ensure_ascii=False, indent=1, default=str)


def build_selected_xlsx(table: SelectedTable, profile_name: str) -> bytes:
    from openpyxl import Workbook

    from app.staging.xlsx_style import write_table_sheet

    wb = Workbook()
    sheet = wb.active
    sheet.title = _sheet_title("Selected Records", [])
    labels = [column.display_label for column in table.columns]
    long_text: list[tuple[str, str, str, str]] = []
    rows = []
    for record_id, values in table.rows:
        # The Long Text sheet names a record by its identifier ("49.102").
        name = str(values.get(table.identifier) or record_id) if table.identifier else record_id
        cells = []
        for column in table.columns:
            spill: list[str] = []
            cells.append(_xlsx_value(values.get(column.canonical_field), spill))
            long_text += [(sheet.title, name, column.display_label, text) for text in spill]
        rows.append(cells)
    wide = {
        label
        for index, label in enumerate(labels)
        if any(isinstance(row[index], str) and len(row[index]) > 80 for row in rows[:200])
    }
    write_table_sheet(sheet, title=table.heading, subtitle=table.subtitle, columns=labels, rows=rows, wide_columns=wide)
    if long_text:
        _write_long_text(wb.create_sheet(title=_sheet_title("Long Text", wb.sheetnames)), profile_name, long_text)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_MEDIA = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "json": "application/json",
}


def build_selected_export(
    workbook: StagingWorkbook, dataset_id: str, manifest: SelectionManifest
) -> tuple[str | bytes, str, str]:
    """(content, media type, filename) for the manifest's format."""

    table = select(workbook, dataset_id, manifest)
    dataset = find_dataset(workbook, dataset_id)
    assert dataset is not None
    if manifest.format == "csv":
        content: str | bytes = build_selected_csv(table)
    elif manifest.format == "json":
        content = build_selected_json(table, workbook, dataset)
    else:
        content = build_selected_xlsx(table, workbook.profile.display_name)
    return content, _MEDIA[manifest.format], _filename(workbook, dataset_id, manifest.format)


# --- Selected fields: label / value lists ----------------------------------------------
# Document-level values (Supplier, Bill-To, Charges & Totals) are shown as
# label / value lists that draw on several datasets. A selection there
# names each value by (dataset, record, field); the export is one
# Section | Field | Value table in the order the reader saw it.


class SelectedField(BaseModel):
    dataset_id: str = Field(min_length=1, max_length=200)
    record_id: str = Field(min_length=1, max_length=500)
    field: str = Field(min_length=1, max_length=500)


class FieldSelectionManifest(BaseModel):
    format: ExportFormat
    fields: list[SelectedField] = Field(min_length=1, max_length=5_000)


_FIELD_COLUMNS = [
    StagingColumn(key="section", canonical_field="section", display_label="Section", value_type="text"),
    StagingColumn(key="field", canonical_field="field", display_label="Field", value_type="text"),
    StagingColumn(key="value", canonical_field="value", display_label="Value", value_type="text"),
]


def _record_label(dataset: StagingDataset, record, column: StagingColumn) -> str:
    """What the value is called: a name / label printed beside it in the
    record (a field list's Field, a charge's label), else the column's."""

    if dataset.cardinality != "single":
        for suffix in (".name", ".label"):
            for key, cell in record.cells.items():
                if key.endswith(suffix) and key != column.canonical_field and cell.value not in (None, ""):
                    return str(cell.value)
    return column.display_label


def select_fields(workbook: StagingWorkbook, manifest: FieldSelectionManifest) -> SelectedTable:
    rows: list[tuple[str, dict[str, object]]] = []
    seen: set[tuple[str, str, str]] = set()
    for item in manifest.fields:
        key = (item.dataset_id, item.record_id, item.field)
        if key in seen:
            continue
        seen.add(key)
        dataset = find_dataset(workbook, item.dataset_id)
        if dataset is None:
            raise SelectionError(f"Unknown dataset {item.dataset_id!r}.")
        column = next((c for c in exportable_columns(dataset) if c.canonical_field == item.field), None)
        if column is None:
            raise SelectionError(f"Field not exportable from {item.dataset_id!r}: {item.field}.")
        record = next((r for r in dataset.records if r.record_id == item.record_id), None)
        if record is None:
            raise SelectionError(f"Record {item.record_id!r} is not in this document's {dataset.display_name}.")
        cell = record.cells.get(item.field)
        rows.append(
            (
                item.record_id,
                {
                    "section": dataset.display_name,
                    "field": _record_label(dataset, record, column),
                    "value": cell.value if cell is not None else None,
                },
            )
        )
    count = len(rows)
    subtitle = f"Extracted from {workbook.document_filename} · {count} selected field{'s' if count != 1 else ''}."
    return SelectedTable("Selected Fields", subtitle, list(_FIELD_COLUMNS), rows, None)


def build_selected_fields_export(
    workbook: StagingWorkbook, manifest: FieldSelectionManifest
) -> tuple[str | bytes, str, str]:
    """(content, media type, filename) for a label / value selection."""

    table = select_fields(workbook, manifest)
    stem = (workbook.document_filename or "export").rsplit(".", 1)[0]
    filename = f"{stem}_selected_fields.{manifest.format}"
    if manifest.format == "csv":
        content: str | bytes = build_selected_csv(table)
    elif manifest.format == "json":
        content = json.dumps(
            {
                "document": workbook.document_filename,
                "fields": [
                    {"section": values["section"], "field": values["field"], "value": values["value"]}
                    for _record_id, values in table.rows
                ],
            },
            ensure_ascii=False,
            indent=1,
            default=str,
        )
    else:
        content = build_selected_xlsx(table, workbook.profile.display_name)
    return content, _MEDIA[manifest.format], filename
