"""The shared professional sheet layout for every Excel export (PDF and
HTML documents alike): a banner title, a one-line subtitle, a spacer row,
the column header row (row 4), then one shaded, wrapped record per row.

Titles, subtitles and column labels are always supplied by the caller from
the document and dataset being exported — this module only styles them."""

from __future__ import annotations

import math
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BANNER_FILL = PatternFill("solid", fgColor="1F3A56")
HEADER_FILL = PatternFill("solid", fgColor="2E5C80")
ROW_FILL = PatternFill("solid", fgColor="DDEFF7")
BORDER = Border(*(Side(style="thin", color="FFFFFF"),) * 4)

HEADER_ROW = 4
# Column widths (characters). Long-text columns are wide; codes stay narrow.
_WIDE = 60
_NORMAL = 22
_NARROW = 10
_LINE_POINTS = 15
# A record shows the start of its long text; the full text is in the cell.
_MAX_ROW_POINTS = 120

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _width(label: str, values: list[object], wide: bool) -> int:
    if wide:
        return _WIDE
    longest = max((len(str(v)) for v in values[:200] if v is not None), default=0)
    natural = max(len(label) + 2, min(longest, 40) + 2)
    if natural <= 12:
        return max(_NARROW, natural)
    return min(max(natural, 14), _NORMAL + 8)


def _lines(value: object, width: int) -> int:
    if value is None:
        return 1
    text = str(value)
    return sum(max(1, math.ceil(len(part) / max(width - 1, 1))) for part in text.split("\n"))


def write_table_sheet(
    ws,
    *,
    title: str,
    subtitle: str | None,
    columns: list[str],
    rows: list[list[object]],
    wide_columns: set[str] = frozenset(),
) -> None:
    """Writes the banner, subtitle, header (row 4) and records into `ws`."""

    count = max(len(columns), 1)
    last = get_column_letter(count)
    ws.append([title])
    ws.append([subtitle])
    ws.append([])
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws["A1"].font = Font(bold=True, size=15, color="FFFFFF")
    ws["A1"].fill = BANNER_FILL
    ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"].font = Font(italic=True, size=9, color="44546A")
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24

    ws.append(columns)
    for cell in ws[HEADER_ROW]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    widths = []
    for index, label in enumerate(columns):
        values = [row[index] for row in rows if index < len(row)]
        width = _width(label, values, label in wide_columns)
        widths.append(width)
        ws.column_dimensions[get_column_letter(index + 1)].width = width
    header_lines = max((_lines(label, w) for label, w in zip(columns, widths)), default=1)
    ws.row_dimensions[HEADER_ROW].height = max(30, header_lines * _LINE_POINTS)

    body = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        ws.append([_ILLEGAL.sub(" ", v) if isinstance(v, str) else v for v in row])
        excel_row = ws.max_row
        for cell in ws[excel_row]:
            cell.fill = ROW_FILL
            cell.alignment = body
            cell.border = BORDER
        lines = max((_lines(v, w) for v, w in zip(row, widths)), default=1)
        ws.row_dimensions[excel_row].height = min(_MAX_ROW_POINTS, max(1, lines) * _LINE_POINTS)

    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{last}{HEADER_ROW + len(rows)}"
