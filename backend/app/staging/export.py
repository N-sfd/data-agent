"""Generic Professional Excel / dataset CSV built from any StagingWorkbook —
used by profiles without a bespoke export (contract_v3 keeps its unchanged
V3 workbook). Every business sheet carries the value plus its page,
evidence and review status, so the export stays source-grounded."""

from __future__ import annotations

import csv
import io
import re

from app.staging.models import StagingDataset, StagingWorkbook

# Excel limits: control characters cannot be stored, and a cell holds at
# most 32,767 characters. CSV/JSON keep the text unchanged.
_XLSX_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_XLSX_CELL_LIMIT = 32_767


def _xlsx_value(value: object, long_text: list[str] | None = None) -> object:
    """A cell value Excel can store. Text beyond the cell limit keeps its
    start and a note; when `long_text` is given the whole text is appended
    to it (written to the Long Text sheet)."""

    if not isinstance(value, str):
        return value
    text = _XLSX_ILLEGAL.sub(" ", value)
    if len(text) > _XLSX_CELL_LIMIT:
        if long_text is not None:
            long_text.append(text)
            where = "the Long Text sheet, the CSV or the JSON export"
        else:
            where = "the CSV or JSON export"
        note = f" … [truncated for Excel: {len(text):,} characters in full — see {where}]"
        text = text[: _XLSX_CELL_LIMIT - len(note)] + note
    return text

# Excel forbids these in sheet names ("Customer / Bill-To") and caps them at 31.
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_title(name: str, taken: list[str]) -> str:
    base = re.sub(r"\s+", " ", _INVALID_SHEET_CHARS.sub("-", name)).strip(" '")[:31] or "Sheet"
    title, n = base, 2
    while title.lower() in (t.lower() for t in taken):
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    return title


def _source_kind(dataset: StagingDataset) -> str | None:
    for record in dataset.records:
        for cell in record.cells.values():
            if cell.provenance:
                return cell.provenance.source_type
    return None


def _location_label(dataset: StagingDataset) -> str:
    # HTML and XML have no pages: records are located by section / element.
    kind = _source_kind(dataset)
    if kind == "html":
        return "Source Section"
    if kind == "xml":
        return "Source Element"
    return "Source Page"


def _location(provenance) -> object:
    if provenance is None:
        return None
    if provenance.source_type == "xml":
        # A record's element (section path [root, record]); a document
        # detail's own element.
        locator = provenance.source_locator
        if locator and len(locator.section_path) > 1:
            return locator.section_path[-1]
        return locator.dom_path if locator else None
    if provenance.source_type == "html":
        locator = provenance.source_locator
        if locator and locator.section_path:
            return " > ".join(locator.section_path)
        return locator.element_id if locator else None
    return provenance.source_page


def _headers(dataset: StagingDataset) -> list[str]:
    headers = [column.display_label for column in dataset.columns]
    if dataset.role == "business":
        headers.append(_location_label(dataset))
        # An XML value is its own evidence; its element path locates it.
        if _source_kind(dataset) != "xml":
            headers.append("Evidence")
        headers.append("Review Status")
    return headers


def _rows(dataset: StagingDataset) -> list[list[object]]:
    xml = _source_kind(dataset) == "xml"
    rows: list[list[object]] = []
    for record in dataset.records:
        cells = [record.cells[c.canonical_field] for c in dataset.columns]
        row: list[object] = [cell.value for cell in cells]
        if dataset.role == "business":
            provenance = next((cell.provenance for cell in cells if cell.provenance), None)
            row.append(_location(provenance))
            if xml is False:
                row.append(provenance.evidence_text if provenance else None)
            row.append(record.record_status)
        rows.append(row)
    return rows


def find_dataset(workbook: StagingWorkbook, dataset_id: str) -> StagingDataset | None:
    return next((d for d in workbook.datasets if d.dataset_id == dataset_id), None)


def build_dataset_csv(dataset: StagingDataset) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(_headers(dataset))
    writer.writerows(_rows(dataset))
    return buffer.getvalue()


_WIDE_LABELS = {"Evidence", "Text", "Description", "Source Section", "Source Element"}
# Text this long reads as a paragraph and gets a wide column.
_WIDE_TEXT = 80
_SOURCE_NAMES = {"pdf": "PDF", "image": "scanned image", "html": "HTML", "xml": "XML", "docx": "Word", "xlsx": "Excel", "text": "text"}


def _wide_columns(headers: list[str], rows: list[list[object]]) -> set[str]:
    wide = {h for h in headers if h in _WIDE_LABELS}
    for index, header in enumerate(headers):
        if any(isinstance(row[index], str) and len(row[index]) > _WIDE_TEXT for row in rows[:200]):
            wide.add(header)
    return wide


def _subtitle(workbook: StagingWorkbook, dataset: StagingDataset) -> str:
    kind = _SOURCE_NAMES.get(_source_kind(dataset) or "")
    source = f"Extracted from {workbook.document_filename}" + (f" ({kind})" if kind else "")
    count = len(dataset.records)
    noun = "record" if count == 1 else "records"
    pieces = [dataset.description, f"{source}. {count} {noun}, one per row."]
    if dataset.role == "business":
        pieces.append("Labels follow the source document; every value carries its source location, evidence and review status.")
    return " ".join(p.strip() for p in pieces if p)


def build_workbook_xlsx(workbook: StagingWorkbook) -> bytes:
    from openpyxl import Workbook

    from app.staging.xlsx_style import write_table_sheet

    profile = workbook.profile.display_name
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    long_text: list[tuple[str, str, str, str]] = []  # sheet, record, field, text
    sheets: list[tuple[str, StagingDataset]] = []

    for dataset in workbook.datasets:
        sheet = wb.create_sheet(title=_sheet_title(dataset.display_name, wb.sheetnames))
        sheets.append((sheet.title, dataset))
        headers = _headers(dataset)
        rows = []
        for record, row in zip(dataset.records, _rows(dataset)):
            cells = []
            for header, value in zip(headers, row):
                spill: list[str] = []
                cells.append(_xlsx_value(value, spill))
                long_text += [(sheet.title, record.record_id, header, text) for text in spill]
            rows.append(cells)
        write_table_sheet(
            sheet,
            title=f"{profile} — {dataset.display_name}",
            subtitle=_subtitle(workbook, dataset),
            columns=headers,
            rows=rows,
            wide_columns=_wide_columns(headers, rows),
        )

    if long_text:
        _write_long_text(wb.create_sheet(title=_sheet_title("Long Text", wb.sheetnames)), profile, long_text)

    about = [
        ["Document", workbook.document_filename],
        ["Profile", f"{profile} ({workbook.profile.profile_id}@{workbook.profile.profile_version})"],
        ["Outcome", f"{workbook.outcome.title} — {workbook.outcome.message}"],
        ["Provenance", "Every business value carries its source location (page for PDF and images, section for HTML, element path for XML) and evidence."],
        ["Review Status", "Verified / Needs Review / Missing — never inferred from OCR confidence."],
        ["Labels", "Column headings keep the document's own terminology (e.g. MAX QUANTITY is never renamed)."],
    ]
    about += [[title, dataset.description or f"{len(dataset.records)} records"] for title, dataset in sheets]
    if long_text:
        about.append(["Long Text", "Full text of any value longer than an Excel cell holds, split into consecutive parts."])
    write_table_sheet(
        readme,
        title=f"{profile} Staging Workbook — {workbook.document_filename}",
        subtitle="What each sheet holds and the rules every value follows.",
        columns=["Item", "Detail"],
        rows=[[_xlsx_value(a), _xlsx_value(b)] for a, b in about],
        wide_columns={"Detail"},
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# Each Long Text part stays well inside Excel's cell limit.
_LONG_TEXT_PART = 30_000


def _write_long_text(sheet, profile: str, entries: list[tuple[str, str, str, str]]) -> None:
    from app.staging.xlsx_style import write_table_sheet

    rows = []
    for sheet_name, record_id, field, text in entries:
        parts = [text[i : i + _LONG_TEXT_PART] for i in range(0, len(text), _LONG_TEXT_PART)]
        rows += [[sheet_name, record_id, field, f"{n} of {len(parts)}", part] for n, part in enumerate(parts, 1)]
    write_table_sheet(
        sheet,
        title=f"{profile} — Long Text",
        subtitle="Values longer than an Excel cell, in full, as consecutive parts.",
        columns=["Sheet", "Record", "Field", "Part", "Text"],
        rows=rows,
        wide_columns={"Text"},
    )


def build_presentation_xlsx(workbook: StagingWorkbook) -> bytes:
    """The workbook a profile's presentation describes: one sheet per
    dataset of its tabs (QA excluded), in tab order, with exactly that
    dataset's columns and complete values. A LONG TEXT sheet is added only
    when a value exceeds an Excel cell."""

    from openpyxl import Workbook

    from app.staging.xlsx_style import write_table_sheet

    profile = workbook.profile.display_name
    by_id = {dataset.dataset_id: dataset for dataset in workbook.datasets}
    wb = Workbook()
    wb.remove(wb.active)
    long_text: list[tuple[str, str, str, str]] = []
    for view in workbook.profile.views:
        for dataset_id in view.dataset_ids:
            dataset = by_id.get(dataset_id)
            if dataset is None or dataset.role == "qa":
                continue
            name = view.label if len(view.dataset_ids) == 1 or dataset.role == "source" else dataset.display_name
            sheet = wb.create_sheet(title=_sheet_title(name.upper(), wb.sheetnames))
            headers = [column.display_label for column in dataset.columns]
            rows = []
            for record in dataset.records:
                row = []
                for column, header in zip(dataset.columns, headers):
                    cell = record.cells.get(column.canonical_field)
                    spill: list[str] = []
                    row.append(_xlsx_value(cell.value if cell else None, spill))
                    long_text += [(sheet.title, record.record_id, header, text) for text in spill]
                rows.append(row)
            write_table_sheet(
                sheet,
                title=f"{profile} — {dataset.display_name}",
                subtitle=" ".join(p for p in (dataset.description, f"Source: {workbook.document_filename}.") if p),
                columns=headers,
                rows=rows,
                wide_columns=_wide_columns(headers, rows),
            )
    if long_text:
        _write_long_text(wb.create_sheet(title="LONG TEXT"), profile, long_text)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- profile workbooks ------------------------------------------------------------------------


def _populated(dataset: StagingDataset) -> bool:
    return any(cell.value not in (None, "") for record in dataset.records for cell in record.cells.values())


def _field_rows(dataset: StagingDataset, with_section: bool) -> list[list[object]]:
    """A single-record dataset as Field | Value (one row per populated field)."""

    rows: list[list[object]] = []
    for record in dataset.records:
        for column in dataset.columns:
            cell = record.cells.get(column.canonical_field)
            if cell is None or cell.value in (None, ""):
                continue
            row = [column.display_label, cell.value]
            rows.append([dataset.display_name, *row] if with_section else row)
    return rows


def _long_rows(dataset: StagingDataset) -> list[list[object]]:
    """Any dataset as Section | Record | Field | Value — used when one sheet
    combines datasets with different columns."""

    rows: list[list[object]] = []
    for index, record in enumerate(dataset.records, start=1):
        for column in dataset.columns:
            cell = record.cells.get(column.canonical_field)
            if cell is None or cell.value in (None, ""):
                continue
            rows.append([dataset.display_name, index, column.display_label, cell.value])
    return rows


def sheet_table(datasets: list[StagingDataset]) -> tuple[list[str], list[list[object]]]:
    """Columns and rows for one sheet. One repeating dataset keeps its own
    columns (the source's headings); single-record datasets read as
    Field | Value; mixed datasets as Section | Record | Field | Value. Values
    are the staged values, unchanged."""

    if len(datasets) == 1 and datasets[0].cardinality == "repeating":
        dataset = datasets[0]
        headers = [column.display_label for column in dataset.columns]
        rows = [
            [record.cells[column.canonical_field].value if column.canonical_field in record.cells else None for column in dataset.columns]
            for record in dataset.records
        ]
        return headers, rows
    if all(dataset.cardinality == "single" for dataset in datasets):
        combined = len(datasets) > 1
        rows = [row for dataset in datasets for row in _field_rows(dataset, combined)]
        return (["Section", "Field", "Value"] if combined else ["Field", "Value"]), rows
    return ["Section", "Record", "Field", "Value"], [row for dataset in datasets for row in _long_rows(dataset)]


def _source_rows(workbook: StagingWorkbook, labels: dict | None) -> list[list[object]]:
    rows: list[list[object]] = [["Document", workbook.document_filename]]
    if labels and labels.get("type_label"):
        rows.append(["Document Type", labels["type_label"]])
    rows += [
        ["Profile", (labels or {}).get("profile_label") or workbook.profile.display_name],
        ["Extraction Outcome", f"{workbook.outcome.title} — {workbook.outcome.message}".strip(" —")],
        ["Records", workbook.qa_summary.record_count],
        ["Values Needing Review", workbook.qa_summary.needs_review],
    ]
    for dataset in workbook.datasets:
        if dataset.role != "source":
            continue
        for record in dataset.records:
            for column in dataset.columns:
                cell = record.cells.get(column.canonical_field)
                if cell is not None and cell.value not in (None, ""):
                    rows.append([column.display_label, cell.value])
    rows.append(["Traceability", "Every value in this workbook is the staged value shown in Data Agent, with its source evidence there."])
    return rows


def build_profile_xlsx(workbook: StagingWorkbook, sheets, labels: dict | None = None) -> bytes:
    """The profile's workbook: its declared sheets, only those with data,
    each in the professional layout (title, description, headings on row
    4). Values are exactly the staged values the Staging Workbook shows."""

    from openpyxl import Workbook

    from app.staging.profile import SOURCE_SHEET
    from app.staging.xlsx_style import write_table_sheet

    profile = (labels or {}).get("profile_label") or workbook.profile.display_name
    by_id = {dataset.dataset_id: dataset for dataset in workbook.datasets}
    wb = Workbook()
    wb.remove(wb.active)
    long_text: list[tuple[str, str, str, str]] = []
    for sheet in sheets:
        if sheet.title == SOURCE_SHEET:
            headers, rows = ["Item", "Detail"], _source_rows(workbook, labels)
            title, description = f"{profile} — Source", "The source document, its resolution and how values trace back to it."
        else:
            datasets = [by_id[i] for i in sheet.dataset_ids if i in by_id and _populated(by_id[i])]
            if sheet.first_populated:
                datasets = datasets[:1]
            if not datasets:
                continue
            headers, rows = sheet_table(datasets)
            if not rows:
                continue
            name = sheet.title or datasets[0].display_name
            title = f"{profile} — {name}"
            description = " ".join(p for p in (datasets[0].description if len(datasets) == 1 else None,) if p)
        ws = wb.create_sheet(title=_sheet_title((sheet.title or name).upper() if sheet.title != SOURCE_SHEET else "SOURCE", wb.sheetnames))
        cells = []
        for index, row in enumerate(rows):
            out = []
            for header, value in zip(headers, row):
                spill: list[str] = []
                out.append(_xlsx_value(value, spill))
                long_text += [(ws.title, str(index + 1), header, text) for text in spill]
            cells.append(out)
        write_table_sheet(
            ws,
            title=title,
            subtitle=" ".join(p for p in (description, f"Source: {workbook.document_filename}.") if p),
            columns=headers,
            rows=cells,
            wide_columns=_wide_columns(headers, cells) | {"Value", "Detail"},
        )
    if long_text:
        _write_long_text(wb.create_sheet(title="LONG TEXT"), profile, long_text)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
